"""기능명세서 엑셀에 정책 90개 체계의 「정책 대응(90개)」 탭을 더한다. 다른 탭은 건드리지 않는다."""
import json, sys, copy
import openpyxl
from openpyxl.styles import Alignment, PatternFill, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

src, rows_path, out = sys.argv[1], sys.argv[2], sys.argv[3]
rows = json.load(open(rows_path, encoding='utf-8'))
wb = openpyxl.load_workbook(src)
old = wb['정책 대응']
head_style = copy.copy(old['A5']._style)
title_font = copy.copy(old['A1'].font)
# 옛 50개 체계 탭(WBS 작업 번호 연결 포함)은 그대로 두고 90개 체계 탭을 바로 뒤에 더한다.
NAME = '정책 대응(90개)'
if NAME in wb.sheetnames:
    wb.remove(wb[NAME])
ws = wb.create_sheet(NAME, wb.sheetnames.index('정책 대응') + 1)

CATS = ['이미 구현', '구현 가능(PaC)', '솔루션 기능으로 커버', '구현 불가·제외']
ws['A1'] = 'MCP 전주기 보안정책 90개 구현 대응 현황'
ws['A1'].font = title_font
ws['A2'] = ('정책 기준: MCP 보안 관리 정책.xlsm(2026-10-27 06:39 수정본) · 코드 기준: main 1d22361 정적 판정(코드 근거, 실기기 집행 확인 아님) · '
            '교차검증: Codex gpt-6.1-sol·gpt-6.1-terra · 작성 2026-10-27')
ws['A3'] = ('분류 — 이미 구현: 시트의 자동 차단 조건을 현재 코드가 집행 · 구현 가능(PaC): OPA/Rego 판정으로 구현 가능(판정 값을 만드는 센서 포함) · '
            '솔루션 기능으로 커버: Rego가 아니라 전송 계층·실행기·저장소·엔드포인트·연동·운영 절차가 집행 · 구현 불가·제외: 관측·집행 불가 또는 현재 범위 제외')
for c in ('A2', 'A3'):
    ws[c].alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells('A2:I2'); ws.merge_cells('A3:I3')
ws.row_dimensions[2].height = 30; ws.row_dimensions[3].height = 45

# 집계(행 5): 분류별 개수
ws['A5'] = '분류'; ws['B5'] = '개수'; ws['C5'] = 'P1'
for c in ('A5', 'B5', 'C5'): ws[c]._style = copy.copy(head_style)
first, last = 10, 9 + len(rows)
for i, cat in enumerate(CATS):
    r = 6 + i
    ws.cell(r, 1, cat)
    ws.cell(r, 2, f'=COUNTIF($E${first}:$E${last},A{r})')
    ws.cell(r, 3, f'=COUNTIFS($E${first}:$E${last},A{r},$D${first}:$D${last},"P1")')

headers = ['정책 ID', '단계', '통제 대상', '우선순위', '구현 분류', '현재 상태', '구현 수단·근거(코드)', '사유(솔루션 커버·불가·제외) / 남은 것', '옛 번호(10-27 05시판)']
for j, h in enumerate(headers, 1):
    c = ws.cell(9, j, h); c._style = copy.copy(head_style)
for i, r in enumerate(rows):
    vals = [r['pol'], r['stage'], r['target'], r['priority'], r['category'], r['status'], r['runtime'] or '—',
            r['reason'] or r['gap'] or '—', r['old'] if r['old'] != r['pol'] else '']
    for j, v in enumerate(vals, 1):
        c = ws.cell(first + i, j, v)
        c.alignment = Alignment(wrap_text=True, vertical='top')

dv = DataValidation(type='list', formula1='"' + ','.join(CATS) + '"', allow_blank=False)
ws.add_data_validation(dv); dv.add(f'E{first}:E{last}')
colors = {'이미 구현': 'D9EAD3', '구현 가능(PaC)': 'CFE2F3', '솔루션 기능으로 커버': 'FFF2CC', '구현 불가·제외': 'F4CCCC'}
for cat, col in colors.items():
    ws.conditional_formatting.add(f'E{first}:E{last}', FormulaRule(formula=[f'$E{first}="{cat}"'], fill=PatternFill('solid', start_color=col, end_color=col)))
for j, w in enumerate([10, 14, 26, 9, 18, 10, 34, 70, 12], 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
ws.freeze_panes = 'B10'
ws.auto_filter.ref = f'A9:I{last}'
wb.save(out)
print('saved', out, len(rows))
