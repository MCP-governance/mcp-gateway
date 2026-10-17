"""Build the compliance-driven development plan from the hand-written plan source and the requirement matrix.

Reads  analysis/compliance/dev-plan/plan-source.json (proposals, acceptance additions, decisions, first sprint; evidence is cited as [doc id, ref token])
       analysis/compliance/requirements.json and pol-legal-basis.json (run build_compliance_matrix.py first)
Writes analysis/compliance/dev-plan/README.md, dev-plan.json and dev-plan.xlsx (regenerated review workbook, not a source of truth).
Evidence that no longer resolves to a provision (e.g. a verifier removed it) is reported and makes the run fail.
Run from the repository root:  python -B -X utf8 research/llm-wiki/tools/build_compliance_plan.py
"""
from __future__ import annotations

import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

WIKI = Path(__file__).resolve().parents[1]
OUT = WIKI / "analysis/compliance"
PLAN = OUT / "dev-plan"
TODAY = "2026-10-17"

THEME_KO = {
    "human_oversight": "인간 감독", "logging_records": "기록·로그", "transparency_labelling": "투명성·표시", "data_protection_lawful_basis": "개인정보 적법 근거",
    "data_minimisation_retention_deletion": "최소화·보유·파기", "cross_border_transfer": "국외 이전", "automated_decision_rights": "자동화된 결정 권리",
    "security_risk_management": "보안 위험 관리", "secure_development_supply_chain": "보안 개발·공급망", "incident_reporting": "사고 보고",
    "access_control_identity": "접근 통제·신원", "third_party_management": "제3자 관리", "risk_impact_assessment": "위험·영향 평가",
    "governance_accountability": "거버넌스·책임", "testing_red_teaming_robustness": "시험·레드티밍·견고성", "agentic_ai": "에이전트 AI",
    "scope_applicability": "적용 범위", "penalties_enforcement": "제재·집행", "other": "기타"}
STRENGTH_KO = {"mandatory": "의무", "recommended": "권고", "informative": "참고"}
STATUS_KO = {"in_force": "시행 중", "voluntary": "자율", "draft": "초안", "adopted_not_yet_applicable": "채택·미시행", "proposed": "제안", "superseded": "대체됨",
             "withdrawn": "철회"}
JUR_KO = {"KR": "한국", "US-federal": "미국 연방", "US-states": "미국 주", "EU": "EU", "CN": "중국", "CA": "캐나다", "SG": "싱가포르", "UK": "영국", "JP": "일본",
          "AU": "호주", "IN": "인도", "INTL": "국제"}


def jur(code: str) -> str:
    if code.startswith("US-") and code not in JUR_KO:
        return "미국 주(" + code[3:] + ")"
    return JUR_KO.get(code, code)


def esc(text) -> str:
    return str(text).replace("|", "\\|").replace("\n", " ")


def load():
    req = json.loads((OUT / "requirements.json").read_text(encoding="utf-8"))
    basis = json.loads((OUT / "pol-legal-basis.json").read_text(encoding="utf-8"))
    src = json.loads((PLAN / "plan-source.json").read_text(encoding="utf-8"))
    return req, basis, src


class Resolver:
    def __init__(self, rows):
        self.by_doc = defaultdict(list)
        for r in rows:
            self.by_doc[r["doc_id"]].append(r)
        self.unresolved = []

    def find(self, doc_id: str, token: str):
        rows = self.by_doc.get(doc_id, [])
        exact = [r for r in rows if r["ref"] == token]
        if exact:
            return exact
        pre = [r for r in rows if r["ref"].startswith(token)]
        if pre:
            return pre
        sub = [r for r in rows if token in r["ref"]]
        if sub:
            return sub
        self.unresolved.append((doc_id, token))
        return []

    def many(self, pairs):
        out, seen = [], set()
        for doc_id, token in pairs:
            for r in self.find(doc_id, token):
                if r["req_id"] not in seen:
                    seen.add(r["req_id"])
                    out.append(r)
        return out


def evidence_row(r) -> dict:
    return {"req_id": r["req_id"], "doc_id": r["doc_id"], "jurisdiction": r["jurisdiction"], "ref": r["ref"], "topic": r["topic_ko"], "strength": r["strength"],
            "binding_doc": r["binding_doc"], "doc_status": r["doc_status"], "timing": r["timing_ko"], "pol_ids": r["pol_ids"]}


def ev_table(rows) -> list[str]:
    L = ["| 요구 ID | 문서 | 관할 | 조항 | 요지 | 강도 | 구속력 | 문서 상태 |", "| --- | --- | --- | --- | --- | --- | --- | --- |"]
    for r in rows:
        L.append(f"| {r['req_id']} | [{r['doc_id']}](../../../papers/{r['doc_id']}.md) | {jur(r['jurisdiction'])} | {esc(r['ref'])} | {esc(r['topic'])} | {STRENGTH_KO[r['strength']]} | "
                 f"{'있음' if r['binding_doc'] else '없음'} | {STATUS_KO.get(r['doc_status'], r['doc_status'])} |")
    return L


def deadline_rows(rows):
    sel = [r for r in rows if r["theme"] == "incident_reporting" and r["timing_ko"] and r["binding_mandatory"]]
    sel.sort(key=lambda r: (r["jurisdiction_group"], r["doc_id"], r["req_id"]))
    return sel


RETENTION = re.compile(r"개월|년간|1년|3년|일 이내|시간 이내|일간")


def retention_rows(rows):
    sel = [r for r in rows if r["theme"] in ("logging_records", "data_minimisation_retention_deletion") and RETENTION.search(r["timing_ko"] or "")]
    sel.sort(key=lambda r: (r["jurisdiction_group"], r["doc_id"], r["req_id"]))
    return sel


def jurisdiction_summary(req):
    docs = {d["id"]: d for d in req["documents"]}
    by = defaultdict(lambda: {"docs": [], "binding_docs": 0, "reqs": 0, "binding_mandatory": 0, "not_yet": []})
    for d in docs.values():
        g = d["jurisdiction_group"]
        by[g]["docs"].append(d["id"])
        by[g]["binding_docs"] += 1 if d["binding"] else 0
        if d["status"] in ("adopted_not_yet_applicable", "draft", "proposed"):
            by[g]["not_yet"].append(d["id"])
    for r in req["requirements"]:
        g = r["jurisdiction_group"]
        by[g]["reqs"] += 1
        by[g]["binding_mandatory"] += 1 if r["binding_mandatory"] else 0
    return by


def main() -> int:
    req, basis, src = load()
    rows = req["requirements"]
    res = Resolver(rows)

    caps = []
    for c in src["capabilities"]:
        ev = res.many(c["evidence"])
        caps.append({**c, "evidence_rows": [evidence_row(r) for r in ev]})
    cap_by = {c["id"]: c for c in caps}

    props = []
    for p in src["proposals"]:
        cap = cap_by[p["capability"]]
        props.append({**p, "evidence_rows": cap["evidence_rows"]})

    adds = []
    for a in src["acceptance_additions"]:
        adds.append({**a, "evidence_rows": [evidence_row(r) for r in res.many(a["evidence"])]})

    deadlines = deadline_rows(rows)
    retention = retention_rows(rows)
    jsum = jurisdiction_summary(req)
    weak = [p for p in basis["policies"] if p["binding_mandatory"] == 0]
    n_docs = len(req["documents"])
    n_req = len(rows)
    n_bm = sum(1 for r in rows if r["binding_mandatory"])

    if res.unresolved:
        print("UNRESOLVED evidence:")
        for doc_id, token in res.unresolved:
            print("  ", doc_id, token)

    # ---- README ----
    L = [f"# 규제 근거 기반 개발 착수 계획 ({TODAY})", "",
         "해외·국내 법·규제·표준 읽기 결과([컴플라이언스 근거 위키](../README.md))를 기능명세·개발 WBS([기능명세서·개발 WBS](../../feature-spec-2026-10-17/README.md))와 연결해, "
         "추가로 개발할 항목, 기존 항목에 더할 수용 기준, 결정이 필요한 사항, 첫 스프린트를 정리한 문서다. 이 문서는 제안이며 구현·승인·법적 적합성 판단이 아니다.", "",
         f"- 입력: 근거 문서 {n_docs}건, 조항 단위 요구사항 {n_req}건(구속력 있는 문서의 '의무' 조항 {n_bm}건). 조항별 근거는 [요구사항 매트릭스](../requirements.json)와 [시트](dev-plan.xlsx).",
         "- 근거 표의 요구 ID는 매트릭스 생성 순서 번호다. 문서 검수에서 조항이 조정되면 번호가 바뀔 수 있어 문서 ID와 조항 표기를 함께 읽는다.",
         "- 모든 제안은 제품 코드·정책·배포를 바꾸지 않은 계획이다. 일정은 팀 확정 전 제안안이다.", ""]

    L += ["## 1. 능력별 규제 요구와 공백", "",
          "| 능력 | 규제가 요구하는 것 | 현재 기능 | 공백 | 제안 | 근거 조항 |", "| --- | --- | --- | --- | --- | ---: |"]
    for c in caps:
        L.append(f"| {c['id']} {esc(c['title'])} | {esc(c['demand'])} | {esc(c['existing'])} | {esc(c['gap'])} | {c['proposal']} | {len(c['evidence_rows'])} |")
    L.append("")

    L += ["## 2. 추가 개발 항목 제안", "",
          "번호는 [백로그](../../../status/backlog.md)의 마지막 항목(DEV-65) 다음 번호를 제안한다. 백로그 반영은 팀 확정 후 한다.", "",
          "| 항목 | 이름 | 능력 | 일정 제안안 | 선행 | 결정 필요 |", "| --- | --- | --- | --- | --- | --- |"]
    for p in props:
        L.append(f"| {p['id']} | {esc(p['title'])} | {p['capability']} | {p['week']} | {', '.join(p['depends']) or '—'} | {', '.join(p['decisions']) or '—'} |")
    L.append("")
    for p in props:
        L += [f"### {p['id']} · {p['title']}", "", f"- 확장 대상: {p['extends']}", f"- 능력: {p['capability']} {cap_by[p['capability']]['title']}", "", "**범위**", ""]
        L += [f"- {x}" for x in p["scope"]]
        L += ["", "**수용 기준**", ""]
        L += [f"- {x}" for x in p["acceptance"]]
        L += ["", f"**위험·한계**: {p['risks']}", "", f"<details><summary>근거 조항 {len(p['evidence_rows'])}건</summary>", ""]
        L += ev_table(p["evidence_rows"])
        L += ["", "</details>", ""]

    L += ["## 3. 기존 개발 항목에 더할 수용 기준", ""]
    for a in adds:
        L += [f"### {a['dev']} · {a['target']}", ""]
        L += [f"- {x}" for x in a["add"]]
        L += ["", f"<details><summary>근거 조항 {len(a['evidence_rows'])}건</summary>", ""]
        L += ev_table(a["evidence_rows"])
        L += ["", "</details>", ""]

    L += ["## 4. 결정이 필요한 사항", "", "| 번호 | 질문 | 선택지 | 권장 | 필요 시점 |", "| --- | --- | --- | --- | --- |"]
    for d in src["decisions"]:
        L.append(f"| {d['id']} | {esc(d['question'])} | {esc(d['options'])} | {esc(d['recommendation'])} | {esc(d['needed_by'])} |")
    L.append("")

    L += ["## 5. 첫 스프린트(2주) 체크리스트", "", "| 번호 | 작업 | 완료 조건 | 관련 항목 |", "| ---: | --- | --- | --- |"]
    for s in src["sprint1"]:
        L.append(f"| {s['n']} | {esc(s['task'])} | {esc(s['done_when'])} | {s['dev']} |")
    L.append("")

    L += ["## 6. 사고 보고 기한 대조표 (구속력 문서의 의무 조항)", "",
          "기한 문구가 있는 구속력 문서의 의무 조항만 모았다. '시행령 위임'이나 '수치 없음'은 그대로 적었다. 이 수치를 관할 팩에 옮길 때는 문서 상태(시행 전·초안)와 최신 개정을 다시 확인한다.", "",
          "| 관할 | 문서 | 조항 | 기한·기간 | 요지 | 문서 상태 |", "| --- | --- | --- | --- | --- | --- |"]
    for r in deadlines:
        L.append(f"| {jur(r['jurisdiction'])} | [{r['doc_id']}](../../../papers/{r['doc_id']}.md) | {esc(r['ref'])} | {esc(r['timing_ko'])} | {esc(r['topic_ko'])} | {STATUS_KO.get(r['doc_status'], r['doc_status'])} |")
    L.append("")

    L += ["## 7. 로그·기록 보존 기간 대조표", "",
          "기록·로그와 보유·파기 주제의 조항 중 기간이 적힌 것이다(권고·참고 포함, 강도 표시). 보존 하한(최소)과 삭제 의무(최대)가 섞여 있어 방향을 조항에서 확인해야 한다.", "",
          "| 관할 | 문서 | 조항 | 기간 | 요지 | 강도 | 구속력 |", "| --- | --- | --- | --- | --- | --- | --- |"]
    for r in retention:
        L.append(f"| {jur(r['jurisdiction'])} | [{r['doc_id']}](../../../papers/{r['doc_id']}.md) | {esc(r['ref'])} | {esc(r['timing_ko'])} | {esc(r['topic_ko'])} | {STRENGTH_KO[r['strength']]} | {'있음' if r['binding_doc'] else '없음'} |")
    L.append("")

    L += ["## 8. 관할별 요약", "", "| 관할 | 문서 | 구속력 문서 | 조항 | 구속력 의무 조항 | 시행 전·초안 문서 |", "| --- | ---: | ---: | ---: | ---: | --- |"]
    for g in sorted(jsum, key=lambda x: (x != "KR", x)):
        s = jsum[g]
        L.append(f"| {jur(g)} | {len(s['docs'])} | {s['binding_docs']} | {s['reqs']} | {s['binding_mandatory']} | {', '.join(s['not_yet']) or '—'} |")
    L.append("")

    L += ["## 9. 법적 근거가 약한 정책", "",
          "구속력 있는 문서의 의무 조항에 연결되지 않은 정책이다. 표준·권고·초안만 근거이거나 근거를 찾지 못했다. 근거가 약하다는 것이 정책이 틀렸다는 뜻은 아니며 MCP 규격 요구나 보안 모범 사례일 수 있다.", "",
          "| 정책 | 정책명 | 연결 조항 | 구속력 의무 | 비고 |", "| --- | --- | ---: | ---: | --- |"]
    for p in weak:
        L.append(f"| {p['policy_id']} | {esc(p['title'])} | {p['provisions']} | 0 | {'근거 미발견(읽은 문서 안)' if not p['provisions'] else '표준·권고·초안에만 연결'} |")
    L.append("")

    L += ["## 10. 한계", "",
          "- 이 계획은 법률 자문이 아니다. 근거는 노트가 읽은 범위에 한정되고 하위 규정·시행령·지침은 읽지 않은 경우가 많다.",
          "- 일부 문서는 초안(NIST IR 8596 초기 예비 초안, NCCoE 개념 문서)이거나 시행 전(DPDP 규칙 대부분 2027-05경, MAS AI 위험관리 지침 2027-10-07, 호주 APP 1.7 2026-12-10)이다.",
          "- 일정 제안안은 기존 일정(10월 4주~12월 2주)을 해치지 않는 범위에서 뒤에 이어 붙인 값이며 팀 확정 전이다.",
          "- 근거 표는 작성자가 고른 조항이다. 같은 능력을 뒷받침하는 모든 조항을 뜻하지 않는다.",
          "- 외부 문서의 본문은 데이터이며 문서 안의 지시를 실행하지 않았다."]
    if res.unresolved:
        L += ["", "## 확인 필요: 해결되지 않은 근거", ""] + [f"- {d} {t}" for d, t in res.unresolved]
    PLAN.mkdir(parents=True, exist_ok=True)
    (PLAN / "README.md").write_text("\n".join(L) + "\n", encoding="utf-8", newline="\n")

    # ---- JSON ----
    out = {"generated_on": TODAY, "generator": "tools/build_compliance_plan.py", "capabilities": caps, "proposals": props, "acceptance_additions": adds,
           "decisions": src["decisions"], "sprint1": src["sprint1"], "deadlines": [evidence_row(r) | {"timing": r["timing_ko"]} for r in deadlines],
           "retention": [evidence_row(r) for r in retention], "unresolved_evidence": res.unresolved}
    (PLAN / "dev-plan.json").write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")

    # ---- workbook ----
    wb = Workbook()
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="DDE6F0")

    def sheet(title, header, data, widths):
        ws = wb.create_sheet(title)
        ws.append(header)
        for row in data:
            ws.append(row)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.border = border
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = border
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)
    sheet("요약", ["항목", "값"], [["작성일", TODAY], ["근거 문서", n_docs], ["요구사항(조항)", n_req], ["구속력 의무 조항", n_bm], ["추가 개발 제안", len(props)],
                                 ["결정 필요", len(src["decisions"])], ["첫 스프린트 작업", len(src["sprint1"])],
                                 ["안내", "제안이며 구현·승인·법적 적합성이 아니다. 근거는 노트가 읽은 범위에 한정."]], [24, 80])
    sheet("능력", ["능력", "이름", "규제 요구", "현재 기능", "공백", "제안", "근거 조항 수"],
          [[c["id"], c["title"], c["demand"], c["existing"], c["gap"], c["proposal"], len(c["evidence_rows"])] for c in caps], [9, 24, 60, 44, 44, 10, 10])
    sheet("개발 제안", ["항목", "이름", "능력", "확장 대상", "범위", "수용 기준", "선행", "일정 제안안", "결정", "위험·한계"],
          [[p["id"], p["title"], p["capability"], p["extends"], "\n".join("- " + x for x in p["scope"]), "\n".join("- " + x for x in p["acceptance"]), ", ".join(p["depends"]),
            p["week"], ", ".join(p["decisions"]), p["risks"]] for p in props], [9, 30, 9, 30, 60, 60, 12, 16, 10, 40])
    sheet("기존 항목 보강", ["항목", "대상", "더할 수용 기준", "근거 조항"],
          [[a["dev"], a["target"], "\n".join("- " + x for x in a["add"]), ", ".join(f"{r['doc_id']} {r['ref']}" for r in a["evidence_rows"])] for a in adds], [9, 36, 70, 60])
    sheet("결정", ["번호", "질문", "선택지", "권장", "필요 시점"], [[d["id"], d["question"], d["options"], d["recommendation"], d["needed_by"]] for d in src["decisions"]], [8, 44, 44, 60, 22])
    sheet("첫 스프린트", ["번호", "작업", "완료 조건", "관련 항목"], [[s["n"], s["task"], s["done_when"], s["dev"]] for s in src["sprint1"]], [6, 70, 50, 14])
    sheet("사고 기한표", ["관할", "문서", "조항", "기한·기간", "요지", "문서 상태", "요구 ID"],
          [[jur(r["jurisdiction"]), r["doc_id"], r["ref"], r["timing_ko"], r["topic_ko"], STATUS_KO.get(r["doc_status"], r["doc_status"]), r["req_id"]] for r in deadlines], [12, 10, 24, 40, 50, 12, 14])
    sheet("보존 기간표", ["관할", "문서", "조항", "기간", "요지", "강도", "구속력", "요구 ID"],
          [[jur(r["jurisdiction"]), r["doc_id"], r["ref"], r["timing_ko"], r["topic_ko"], STRENGTH_KO[r["strength"]], "있음" if r["binding_doc"] else "없음", r["req_id"]] for r in retention],
          [12, 10, 24, 40, 50, 8, 8, 14])
    sheet("관할 요약", ["관할", "문서", "구속력 문서", "조항", "구속력 의무 조항", "시행 전·초안 문서"],
          [[jur(g), len(s["docs"]), s["binding_docs"], s["reqs"], s["binding_mandatory"], ", ".join(s["not_yet"])] for g, s in sorted(jsum.items(), key=lambda x: (x[0] != "KR", x[0]))],
          [12, 8, 12, 8, 14, 40])
    sheet("근거 조항", ["제안·능력", "요구 ID", "문서", "관할", "조항", "요지", "강도", "구속력", "문서 상태"],
          [[c["id"], r["req_id"], r["doc_id"], jur(r["jurisdiction"]), r["ref"], r["topic"], STRENGTH_KO[r["strength"]], "있음" if r["binding_doc"] else "없음",
            STATUS_KO.get(r["doc_status"], r["doc_status"])] for c in caps for r in c["evidence_rows"]], [10, 14, 10, 10, 26, 60, 8, 8, 12])
    wb.save(PLAN / "dev-plan.xlsx")
    print(f"plan: {len(caps)} capabilities, {len(props)} proposals, {len(adds)} additions, {len(deadlines)} deadlines, {len(retention)} retention rows, unresolved {len(res.unresolved)}")
    return 1 if res.unresolved else 0


if __name__ == "__main__":
    sys.exit(main())
