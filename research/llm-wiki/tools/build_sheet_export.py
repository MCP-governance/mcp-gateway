"""Export the unified sheet locally; never execute or rewrite source workbooks.

Python 3.12, standard library and openpyxl. Run with -B -X utf8.
"""
from __future__ import annotations

import argparse
import ast
import csv
import hashlib
import io
import json
import posixpath
import re
import subprocess
import tempfile
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from urllib.parse import quote, unquote

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
WIKI = ROOT / "research/llm-wiki"
OUT = WIKI / "sheet"
GITHUB = "https://github.com/MCP-governance/mcp-gateway/blob/main/"
RAW = "https://raw.githubusercontent.com/MCP-governance/mcp-gateway/<커밋>/research/llm-wiki/sheet/tabs/"
SEP = " · "
MAX_CELL = 45_000
PATTERNS = {
    "wbs": r"G\d{2}\.\d{2}", "features": r"F-\d{2}",
    "policies": r"POL-\d{2}", "backlog": r"(?:DEV|WIKI)-\d{2}",
    "decisions": r"D-\d+", "experiments": r"E\d{2}",
    "components": r"CMP-\d{2}", "oss-reuse": r"OSS-\d{2}", "docs": r"DOC-\d{3,}",
    "deploys": r"DEP-\d{2,}", "commits": r"[0-9a-f]{7}",
}
TOKEN = re.compile(r"(?<![A-Za-z0-9_.-])(?:" + "|".join(PATTERNS.values()) + r")(?![A-Za-z0-9_.-])")
TABS = [
    (1, "요약", None, "관리", None), (2, "WBS", "wbs", "관리", "관리"),
    (3, "백로그", "backlog", "관리", "관리"), (4, "결정", "decisions", "관리", "관리"),
    (5, "개발이력", "commits", "관리", "참조"), (6, "배포", "deploys", "관리", "관리"),
    (7, "기능", "features", "제품", "관리"), (8, "구성요소", "components", "제품", "관리"),
    (9, "MVP 기능", "mvp", "제품", "원본"), (10, "원 계획", "plan-v15", "관리", "원본"),
    (11, "정책", "policies", "정책", "관리"), (12, "정책 상세", "policies-detail", "정책", "원본"),
    (13, "정책 독스용", "policies-docs", "정책", "원본"), (14, "추가 정책", "policies-extra", "정책", "원본"),
    (15, "정책 번호 대응", "policy-id-map", "정책", "참조"), (16, "시험", "experiments", "검증", "관리"),
    (17, "위협", "threats", "검증", "참조"), (18, "로컬 도구 위험", "local-tool-risks", "검증", "참조"),
    (19, "권한 근거", "endpoint-permission-refs", "검증", "참조"), (20, "규제 문서", "reg-docs", "규제", "참조"),
    (21, "규제 요구", "reg-requirements", "규제", "참조"), (22, "주제×관할", "reg-matrix", "규제", "참조"),
    (23, "규제 계획 요약", "reg-plan-summary", "규제", "참조"), (24, "규제 능력", "reg-capabilities", "규제", "참조"),
    (25, "규제 개발 제안", "reg-proposals", "규제", "참조"), (26, "규제 보강", "reg-reinforcements", "규제", "참조"),
    (27, "규제 결정 대기", "reg-decisions", "규제", "참조"), (28, "규제 첫 스프린트", "reg-sprint", "규제", "참조"),
    (29, "사고 기한", "reg-incident-deadlines", "규제", "참조"), (30, "보존 기간", "reg-retention", "규제", "참조"),
    (31, "관할 요약", "reg-jurisdictions", "규제", "참조"), (32, "규제 근거 조항", "reg-clauses", "규제", "참조"),
    (33, "기업 요구", "ent-requirements", "기업", "참조"), (34, "코드 구조", "code-structure", "기술", "참조"),
    (35, "제품 기능 대조", "competitor-map", "기업", "참조"), (36, "연동 대상", "integrations", "기업", "참조"),
    (37, "연동 시험", "integration-cases", "기업", "참조"), (38, "MCP 대상", "mcp-candidates", "기업", "참조"),
    (39, "스킬 후보", "skill-candidates", "기술", "참조"), (40, "OSS 재사용", "oss-reuse", "기술", "참조"),
    (41, "제품 분석", "solutions", "기업", "참조"), (42, "논문", "papers", "기술", "참조"),
    (43, "레퍼런스 갱신", "reference-updates", "기술", "참조"), (44, "도입 기록", "imports", "기술", "참조"),
    (45, "자료", "docs", "관리", "관리"),
]
HEADERS = {
    "wbs": "ID|목표|목표명|작업|구분|우선|구현|회귀|배포|실기기|검수|진척률|상태|담당|진행 메모|계획 일정|완료 조건|남은 일|기능|정책|백로그|결정|시험|구성요소|코드|회귀 시험|근거 자료",
    "backlog": "ID|작업|상태|진행|완료 조건|WBS|결정|근거 자료|WBS 상태",
    "decisions": "ID|날짜|제목|요지|백로그|WBS|커밋|문서",
    "commits": "커밋|날짜|작성자|유형|범위|요약|결정|백로그|WBS|병합|기기 배포|링크",
    "deploys": "ID|날짜|대상|커밋|브랜치|내용|근거 자료",
    "features": "ID|대분류|기능|설명|상태|정책|구현 수단|비고|WBS|연결 근거|WBS 상태",
    "components": "ID|이름|분류|상태|위치|역할|버전|라이선스|저장소|WBS|OSS 재사용|구조도|아이콘",
    "policies": "ID|영역|정책명|정책 내용|CIA 영향|자동화 분류|현재|개발 후|관련 추가 개발|자동화 한계|기능|WBS|규제 조항 수|구속력 의무|요구 ID|번호 대응|근거|WBS 작업 수|WBS 평균 진척|기능 수",
    "experiments": "ID|이름|목적|조건|판정 기준|상태|정책|WBS|백로그|카드",
    "docs": "ID|제목|종류|날짜|경로|관련 ID|상태",
}
# 시트에서 수식으로 채우는 열(저장소 시드에서는 빈 값)
FORMULA_COLUMNS = {"wbs": {"진척률", "상태"}, "backlog": {"WBS 상태"}, "features": {"WBS 상태"},
                   "policies": {"WBS 작업 수", "WBS 평균 진척", "기능 수"}}
LABELS = dict(x.split(":", 1) for x in (
    "title:이름|name:이름|category:구분|kind:종류|area:영역|requirement:要求|acceptance:완료 조건|"
    "negative_cases:음성·경계|status:상태|priority:우선|gap:남은 일|code_evidence:코드|test_evidence:시험 근거|"
    "source_refs:출처|requirement_id:요구 ID|product:제품|version_or_commit:판본|feature:기능|"
    "evidence_url:근거 링크|evidence_paths:근거 자료|verification:확인 수준|our_gap:남은 일|analysis_page:분석 문서|"
    "route:경로|auth:인증|capabilities:지원 기능|license:라이선스|compatibility_status:연동 상태|source_url:출처|"
    "prerequisites:조건|procedure:절차|expected:판정 기준|negative_case:음성 입력|evidence:근거|failure_fix:실패 시 수정|"
    "implementation_status:구현 상태|current_gap:남은 일|package:패키지|version:버전|transport:전송|"
    "auth_boundary:인증 경계|experiment_ids:시험|execution_status:실행 상태|limitations:한계|"
    "repository:저장소|commit:커밋|reuse_plan:활용 방법|unreviewed_scope:미확인 범위|existing_source_path:기존 코드|"
    "implementation_guidance:활용 방법|selection_gate:선정 조건|path:문서|url:링크|summary:요약|verified_on:확인일|"
    "threat:위협|boundary:경계|source_id:원본 ID|runtime_candidates:정책 후보|code_candidates:코드 후보|"
    "local_tool_risk_ids:위험 ID|tools:도구|observation:관측|source_controls:원본 통제|required_control:필요 통제|remaining:남은 일|"
    "read_scope:읽은 범위|official_claim:공식 내용|application:적용|validation_needed:검증 조건|"
    "source_project:원천 프로젝트|upstream_id:원천 ID|imported_on:도입일|files:파일|distribution:배포 조건|provenance:출처 기록|"
    "doi:DOI|publication_status:게재 상태|venue:학술지·학회|published_on:게재일|reading_scope:읽은 범위|"
    "methodology_adaptation:적용 방법|adoption_status:도입 상태|existing_catalog_id:기존 자료|"
    "jurisdiction:관할|instrument_type:문서 유형|binding:구속력|key_dates:주요 일자|official_page_url:공식 페이지|"
    "provisions:조항 수|req_id:요구 ID|doc_id:규제 문서|ref:조항|topic_ko:주제|obligation_ko:요구 내용|"
    "actor:대상|strength:강도|theme:주제 분류|timing_ko:기한|gateway_control_ko:통제|pol_ids:정책|"
    "binding_mandatory:구속력 의무|demand:규제 요구|existing:현재 기능|proposal:제안|capability:능력|"
    "extends:확장 대상|scope:범위|depends:선행|week:일정|decisions:결정 대기|risks:위험·한계|"
    "dev:백로그|target:대상|add:추가 조건|question:질문|options:선택지|recommendation:권장|needed_by:필요 시점|"
    "task:작업|done_when:완료 조건|timing:기한|topic:주제|doc_status:문서 상태|binding_doc:구속력|"
    "authors:저자|year:연도|engineering_use:관련 작업|page:문서|edition:판본|license_url:라이선스 링크|"
    "originals:원본|extracts:추출본|supplier:공급자|maintainer:유지보수|obtained:확보일"
).replace("要求", "요구").split("|") if x)
VALUES = {"mandatory": "의무", "recommended": "권고", "informative": "참고", "not_run": "미실행",
          "in_force": "시행 중", "voluntary": "자율", "draft": "초안", "proposed": "제안",
          "adopted_not_yet_applicable": "채택·미시행", "superseded": "대체됨", "withdrawn": "철회"}
JUR = {"KR": "한국", "AU": "호주", "CA": "캐나다", "CN": "중국", "EU": "EU", "IN": "인도",
       "INTL": "국제", "JP": "일본", "SG": "싱가포르", "UK": "영국", "US-federal": "미국 연방", "US-states": "미국 주"}


def git(*args):
    return subprocess.run(["git", "-C", str(ROOT), *args], check=True, capture_output=True,
                          encoding="utf-8").stdout.strip()


def read(path):
    return Path(path).read_text(encoding="utf-8-sig")


def load(path):
    return json.loads(read(path))


def dump(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def unique(values):
    return list(dict.fromkeys(str(x) for x in values if x is not None and str(x)))


def joined(values):
    return SEP.join(unique(values))


def explicit_list(value, prefix):
    """Expand a source's explicit shorthand, retaining its prefix and range bounds."""
    text = str(value or "")
    found = []
    for match in re.finditer(re.escape(prefix) + r"(\d+)((?:\s*[·,~]\s*\d+)*)", text):
        current = int(match[1])
        found.append(f"{prefix}{current:02d}")
        for tail in re.finditer(r"([·,~])\s*(\d+)", match[2]):
            number = int(tail[2])
            if tail[1] == "~":
                if number < current or number - current > 1000:
                    raise ValueError(f"ID 범위 오류: {value}")
                found.extend(f"{prefix}{n:02d}" for n in range(current + 1, number + 1))
            else:
                found.append(f"{prefix}{number:02d}")
            current = number
    return joined(found)


def ids(value, kind=None):
    text = dump(value) if isinstance(value, (dict, list)) else str(value or "")
    found = unique(m.group() for m in TOKEN.finditer(text))
    return [s for s in found if not kind or re.fullmatch(PATTERNS[kind], s)]


def reverse(rows, column):
    result = defaultdict(list)
    for row in rows:
        for key in str(row.get(column, "")).split(SEP):
            if key:
                result[key].append(row["ID"])
    return {k: joined(v) for k, v in result.items()}


def repo_path(path):
    return Path(path).resolve().relative_to(ROOT).as_posix()


def url(path, base=None):
    path = str(path).strip().replace("\\", "/")
    if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", path) and not re.match(r"^[A-Za-z]:/", path):
        return path
    suffix = ""
    if "::" in path:
        path, symbol = path.split("::", 1)
        suffix = "::" + symbol
    match = re.match(r"^(.*?):(\d+)(?:[-:](\d+))?(.*)$", path)
    fragment = ""
    if match:
        path, start, end, tail = match.groups()
        fragment = f"#L{start}" + (f"-L{end}" if end else "")
        suffix = tail + suffix
    else:
        path, mark, frag = path.partition("#")
        fragment = "#" + frag if mark else ""
    path = unquote(path)
    if base is not None:
        path = posixpath.join(repo_path(base.parent), path)
    else:
        first = path.split("/", 1)[0]
        if first in {"analysis", "catalogs", "experiments", "papers", "sources", "solutions", "threats", "methodology", "status", "integrations"}:
            path = "research/llm-wiki/" + path
    path = posixpath.normpath(path).lstrip("/")
    if path.startswith("../"):
        return str(path) + fragment + suffix
    return GITHUB + quote(path, safe="/") + fragment + suffix


def markdown_links(text, source):
    return re.sub(r"(!?\[[^\]]*\]\()([^\s)]+)([^)]*\))",
                  lambda m: m[1] + url(m[2], source) + m[3], str(text))


def linked_paths(paths):
    return joined(url(p) for p in paths)


def path_links(value):
    if isinstance(value, dict):
        return {k: path_links(v) for k, v in value.items()}
    if isinstance(value, list):
        return [path_links(v) for v in value]
    if not isinstance(value, str):
        return value
    pattern = r"(?<![\w/])(?:research|docs|full_stack_lab|endpoint-agent|analysis|catalogs|experiments|papers|sources|solutions|threats|methodology|status|integrations)/[^\s`<>\[\]{}\"',;|)]+"
    return re.sub(pattern, lambda m: url(m[0]), value)


def cell(value, flat=False):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "있음" if value else "없음"
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (date, datetime)):
        value = value.isoformat()
    elif isinstance(value, list):
        value = joined(value) if all(not isinstance(v, (dict, list)) for v in value) else dump(value)
    elif isinstance(value, dict):
        value = dump(value)
    value = str(value).replace("\r\n", "\n").replace("\r", "\n")
    if flat:
        value = value.replace("\n", " / ")
    # Whitespace-prefixed formula markers are neutralised too, without shifting IDs.
    value = re.sub(r"^(\s*)([=+@])", lambda m: m[1] + {"=": "＝", "+": "＋", "@": "＠"}[m[2]], value)
    return value[:MAX_CELL - 1] + "…" if len(value) > MAX_CELL else value


def csv_bytes(headers, rows, flat=False):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([cell(row.get(h) if h == "ID" else path_links(row.get(h)), flat) for h in headers])
    return stream.getvalue().encode("utf-8")


def md_tables(text):
    lines = text.splitlines()
    for index, line in enumerate(lines[:-1]):
        if not line.startswith("|") or not re.match(r"^\|[\s:|-]+\|\s*$", lines[index + 1]):
            continue
        split = lambda s: [x.strip().replace(r"\|", "|") for x in re.split(r"(?<!\\)\|", s.strip().strip("|"))]
        headers = split(line)
        rows = []
        for body in lines[index + 2:]:
            if not body.startswith("|"):
                break
            rows.append(dict(zip(headers, split(body))))
        yield headers, rows


def workbook_table(path, sheet, *, plan=False):
    book = load_workbook(path, read_only=True, data_only=True, keep_vba=False)
    try:
        values = list(book[sheet].values)
    finally:
        book.close()
    header_index = next((i for i, row in enumerate(values[:30])
                         if (plan and row and row[0] == "구분") or
                         (not plan and sum(v not in (None, "") for v in row) >= 4 and
                          any(str(v).strip() in {"ID", "정책 ID", "정책 번호", "대분류", "요구 ID", "능력", "항목", "번호", "파일", "관할", "주제", "제안·능력", "검증 층"} for v in row))), None)
    if header_index is None:
        raise ValueError(f"머리글을 찾을 수 없음: {path} / {sheet}")
    raw = values[header_index]
    if plan:
        raw = ("구분", "대분류", "작업", "논문", "산출물", *[f"{month}월 {week}주" for month, weeks in [(9, 4), (10, 5), (11, 4), (12, 4)] for week in range(1, weeks + 1)])
        header_index += 1
    names, count = [], Counter()
    for i, value in enumerate(raw):
        label = str(value or f"열 {i + 1}").replace("\r", "").replace("\n", " ").strip()
        label = {"commit": "커밋", "async": "비동기", "JSON 정본": "원천", "정책 ID": "ID", "정책 번호": "ID"}.get(label, label)
        for start, short in [("MCP Specification", "MCP 명세"), ("MCP Security Best Practices", "MCP 보안 권고"),
                             ("OWASP MCP Top 10", "OWASP 상위 위협"), ("OWASP MCP Security Cheat Sheet", "OWASP 보안 지침"),
                             ("ISO 42001", "ISO 42001"), ("ISMS-P", "ISMS-P"), ("NIST AI RMF", "NIST 위험관리"),
                             ("개인정보 보호법", "개인정보 보호법"), ("인공지능기본법", "인공지능기본법")]:
            if label.startswith(start):
                label = short
                break
        count[label] += 1
        names.append(label if count[label] == 1 else f"{label} {count[label]}")
    result = []
    carried = {}
    for n, values_row in enumerate(values[header_index + 1:], header_index + 2):
        if not any(v is not None and str(v).strip() for v in values_row):
            continue
        row = dict(zip(names, values_row))
        if "ID" in names and not re.fullmatch(r"(?:POL-\d+|PAC\d+|[A-Z][A-Z0-9-]*\d+)", str(row.get("ID", ""))):
            continue
        for key in ("영역", "대분류", "구분"):
            if key in row:
                if row[key] not in (None, ""):
                    carried[key] = row[key]
                elif key in carried:
                    row[key] = carried[key]
        row["_row"] = n
        result.append(row)
    return names, result


def backlog_state(progress):
    if re.search(r"보류|중단|폐기|범위 밖|사용자 결정|결정 대기", progress):
        return "보류"
    if re.search(r"미착수|착수 전", progress):
        return "미착수"
    if re.search(r"구현|수정|배포|실기기|통과|반영|설치", progress) and not re.search(r"미구현|미완료|미실행|미배포|남음|남았|필요|후속|대기|일부", progress):
        return "구현 완료(검수 전)"
    if re.search(r"초안|설계|조사|검토|계획|정의|문서", progress) and not re.search(r"구현|수정|배포|실기기|통과|설치", progress):
        return "설계·조사"
    return "진행 중"


def anchor(heading):
    heading = re.sub(r"<[^>]*>", "", heading).lower().replace("`", "")
    return "".join(c for c in heading if c in " -_" or unicodedata.category(c)[0] in "LN").replace(" ", "-")


def filename_date(name):
    for match in re.finditer(r"(?<!\d)(20\d{2})([-_]?)(\d{2})\2(\d{2})(?!\d)", name):
        try:
            return date(int(match[1]), int(match[3]), int(match[4])).isoformat()
        except ValueError:
            continue
    return ""


def replaced_document(text):
    return bool(re.search(
        r"(?:이 문서|본 문서|이 페이지|이 파일)(?:는|은|가|이| 전체)?\s*[^\n。]{0,80}?"
        r"(?:[으]?로 대체(?:됐|되었|됨|되었음)|더 이상 사용하지 않)|"
        r"^>\s*(?:대체됨|Superseded)\b|"
        r"이전 (?:계획 )?기록[^\n]*\n[^\n]*현재 정본은",
        text, re.I | re.M))


class Export:
    def __init__(self):
        self.tables = {}
        self.notes = []
        self.registry = load(OUT / "manifest.json").get("id_registry", {}) if (OUT / "manifest.json").exists() else {}
        self.decimal_decisions = {}

    def stable_ids(self, slug, keys, prefix, width):
        mapping = self.registry.setdefault(slug, {})
        high = max([int(v.removeprefix(prefix)) for v in mapping.values()] or [0])
        for key in sorted(set(keys)):
            if key not in mapping:
                high += 1
                mapping[key] = f"{prefix}{high:0{width}d}"
        return mapping

    def add(self, slug, rows, sources, headers=None, links=None):
        self.tables[slug] = {"rows": rows, "headers": headers or HEADERS[slug].split("|"),
                             "sources": unique(repo_path(p) if isinstance(p, Path) else p for p in sources), "links": links or {}}
        return rows

    def decision_ids(self, value):
        return unique(self.decimal_decisions.get(int(x[2:]), x) for x in ids(value, "decisions"))

    def generic(self, slug, data, source, fields, *, id_key="id", links=None, details=True):
        if isinstance(data, dict):
            data = data["records"]
        headers = ["ID"] + [LABELS.get(k, k) for k in fields]
        if len(headers) != len(set(headers)):
            raise ValueError(f"중복 열 이름: {slug}: {headers}")
        rows = []
        for record in data:
            row = {"ID": record[id_key]}
            for key in fields:
                value = record.get(key)
                if isinstance(value, str):
                    if key == "jurisdiction":
                        value = JUR.get(value, value)
                    elif key in {"path", "page", "analysis_page", "existing_source_path", "provenance"} and value:
                        value = url(value)
                    else:
                        value = VALUES.get(value, value)
                row[LABELS.get(key, key)] = value
            extra = {k: v for k, v in record.items() if k not in [id_key, *fields]}
            if details:
                row["세부"] = extra if extra else ""
            rows.append(row)
        if details:
            headers.append("세부")
        return self.add(slug, rows, [source], headers, links)

    def management(self):
        decision_path = ROOT / "docs/ai/DECISIONS.md"
        text = read(decision_path)
        matches = list(re.finditer(r"^## (D-\d+)\s+(.+)$", text, re.M))
        self.decimal_decisions = {int(m[1][2:]): m[1] for m in matches}
        decisions = []
        seen_anchors = Counter()
        for i, match in enumerate(matches):
            body = text[match.end():matches[i + 1].start() if i + 1 < len(matches) else len(text)].strip()
            heading = match[1] + " " + match[2]
            a = anchor(heading)
            seen_anchors[a] += 1
            a += f"-{seen_anchors[a] - 1}" if seen_anchors[a] > 1 else ""
            dates = re.findall(r"20\d{2}-\d{2}-\d{2}", match[2])
            plain = re.sub(r"\*\*|`", "", body)
            plain = re.sub(r"^[-*>]\s*", "", plain, flags=re.M)
            sentences = re.split(r"(?<=[.!?。])(?:\s+|$)", plain)
            summary = " ".join(sentences[:2]).replace("\n", " ")
            decisions.append({"ID": match[1], "날짜": dates[-1] if dates else "", "제목": match[2],
                              "요지": summary[:199] + "…" if len(summary) > 200 else summary,
                              "백로그": joined(ids(body, "backlog")), "WBS": joined(ids(body, "wbs")),
                              "커밋": "", "문서": url(repo_path(decision_path)) + "#" + quote(a, safe="-_")})
        self.add("decisions", decisions, [decision_path], links={"백로그": "backlog", "WBS": "wbs", "커밋": "commits"})

        backlog_path = WIKI / "status/backlog.md"
        backlog = []
        for headers, rows in md_tables(read(backlog_path)):
            if "ID" not in headers:
                continue
            for r in rows:
                original_id = r["ID"]
                if original_id == "RES-01":
                    aliases = self.registry.setdefault("backlog_aliases", {"RES-01": "WIKI-02"})
                    r["ID"] = aliases[original_id]
                elif not re.fullmatch(PATTERNS["backlog"], original_id):
                    raise ValueError(f"백로그 원본 ID 형식 확인 필요: {original_id}")
                progress = r[headers[-1]]
                backlog.append({"ID": r["ID"], "작업": markdown_links(r.get("작업", ""), backlog_path),
                                "상태": backlog_state(progress), "진행": markdown_links(progress, backlog_path),
                                "완료 조건": markdown_links(r.get("완료 조건", ""), backlog_path),
                                "WBS": joined(ids(r, "wbs")), "결정": joined(self.decision_ids(r)),
                                "근거 자료": markdown_links(r.get("근거", ""), backlog_path) + (f" · 원본 ID: {original_id}" if original_id != r["ID"] else "")})
        self.add("backlog", backlog, [backlog_path], links={"WBS": "wbs", "결정": "decisions"})

        snap_path = WIKI / "analysis/development-wbs/snapshot.json"
        spec_path = WIKI / "analysis/feature-spec-2026-10-17/feature-spec.json"
        owners_path = ROOT / "research/development-status-site/owners.json"
        snap, spec, owners = load(snap_path), load(spec_path), load(owners_path)
        goal_names = {r["id"]: r["title"] for r in snap["goals"]}
        specs = {r["id"]: r for r in spec["wbs"]}
        snapshot = {r["id"]: r for r in snap["items"]}
        wbs = []
        for ident in unique([*snapshot, *specs]):
            r, s = snapshot.get(ident, {}), specs.get(ident, {})
            goal = r.get("goal_id", ident.split(".")[0])
            row = {"ID": ident, "목표": goal, "목표명": goal_names.get(goal, ""), "작업": r.get("title", s.get("task", "")),
                   "구분": s.get("kind", r.get("kind", "현재 개발")), "우선": r.get("priority", ""),
                   "담당": owners.get(ident, {}).get("owner", ""), "진행 메모": owners.get(ident, {}).get("memo", ""),
                   "계획 일정": s.get("planned_weeks") or r.get("planned_weeks", ""),
                   "완료 조건": r.get("completion_condition", s.get("completion", "")),
                   "남은 일": r.get("known_gap", s.get("remaining", "")), "정책": joined(ids(s.get("policies", r.get("pol_ids", [])), "policies")),
                   "백로그": joined(ids(r, "backlog")), "결정": joined(self.decision_ids(r)), "시험": joined(ids(r, "experiments")),
                   "코드": linked_paths(r.get("code_symbols_files", [])), "회귀 시험": linked_paths(r.get("regression_test_candidates", [])),
                   "근거 자료": linked_paths(r.get("actual_evidence_paths", []) + [e["path"] for e in r.get("evidence", []) if "path" in e])}
            scores = r.get("stage_scores", {})
            if not scores:
                raise ValueError(f"원본 점수 누락: {ident}")
            for key, label in zip(["implementation", "regression", "deployment", "real_device", "acceptance_complete"], ["구현", "회귀", "배포", "실기기", "검수"]):
                row[label] = scores[key]
            wbs.append(row)
        self.add("wbs", wbs, [snap_path, spec_path, owners_path], links={"기능": "features", "정책": "policies", "백로그": "backlog", "결정": "decisions", "시험": "experiments", "구성요소": "components"})

        seed = OUT / "seed/components.csv"
        oss_headers, oss = HEADERS["components"].split("|"), []
        if seed.exists():
            raw = seed.read_bytes().replace(b"\r\n", b"\n")  # Windows 체크아웃(CRLF)도 같은 바이트로
            if raw.startswith(b"\xef\xbb\xbf") or b"\r" in raw:
                raise ValueError("OSS 원본 CSV는 BOM 없이 LF여야 함")
            reader = csv.DictReader(io.StringIO(raw.decode("utf-8"), newline=""))
            oss_headers = reader.fieldnames
            if not oss_headers or oss_headers[0] != "ID" or len(oss_headers) != len(set(oss_headers)):
                raise ValueError("OSS 원본 CSV 머리글 오류")
            oss = list(reader)
            if any(None in r or None in r.values() for r in oss):
                raise ValueError("OSS 원본 CSV 열 수 오류")
            for row in oss:
                if any(cell(v) != v for v in row.values()):
                    raise ValueError(f"OSS 원본 CSV 셀 형식 오류: {row['ID']}")
        self.add("components", oss, [seed] if seed.exists() else [], oss_headers,
                 {k: v for k, v in {"WBS": "wbs", "OSS 재사용": "oss-reuse"}.items() if k in oss_headers})

        review_path = spec_path.parent / "README.md"
        review_rows = [r for _, rows in md_tables(read(review_path)) for r in rows if "리뷰" in r]
        features = []
        seed_path = OUT / "seed/feature-wbs.csv"
        seed_map = {r["ID"]: r["WBS"] for r in csv.DictReader(io.StringIO(read(seed_path)))} if seed_path.exists() else {}
        common = {"사용자", "게이트웨이", "엔드포인트", "서버", "도구", "기능", "정책", "관리", "확인", "등록", "검증", "제한", "처리", "기록", "범위", "호출"}
        for n, r in enumerate(spec["feature_spec"], 1):
            explicit = ids(r, "wbs")
            for review in review_rows:
                label = review["리뷰"]
                if r["feature"] == label or (len(r["feature"]) >= 6 and r["feature"] in label):
                    explicit += ids(review, "wbs")
            explicit = unique(explicit)
            linked = list(explicit)
            evidence = [f"{x}: 명시" for x in explicit]
            tokens = {x for x in re.findall(r"[가-힣A-Za-z0-9]+", r["feature"]) if len(x) >= 2 and x not in common}
            for task in wbs:
                if task["ID"] in linked or not set(ids(task["정책"], "policies")) & set(ids(r["policies"], "policies")):
                    continue
                matches = {x for x in tokens if x in task["작업"]}
                if r["feature"] in task["작업"] or len(matches) >= 2 or any(len(x) >= 4 for x in matches):
                    linked.append(task["ID"])
                    evidence.append(f"{task['ID']}: 정책 일치")
            for extra in ids(seed_map.get(f"F-{n:02d}", ""), "wbs"):
                if extra not in linked:
                    linked.append(extra)
                    evidence.append(f"{extra}: 검토 매핑")
            features.append({"ID": f"F-{n:02d}", "대분류": r["category"], "기능": r["feature"], "설명": r["description"],
                             "상태": r["status"], "정책": joined(ids(r["policies"], "policies")), "구현 수단": r["means"],
                             "비고": r["note"], "WBS": joined(linked), "연결 근거": joined(evidence)})
        self.add("features", features, [spec_path, review_path, seed_path], links={"정책": "policies", "WBS": "wbs"})
        feature_reverse, oss_reverse, backlog_reverse = reverse(features, "WBS"), reverse(oss, "WBS"), reverse(wbs, "백로그")
        for row in wbs:
            row["기능"], row["구성요소"] = feature_reverse.get(row["ID"], ""), oss_reverse.get(row["ID"], "")
        for row in backlog:
            row["WBS"] = joined([*ids(row["WBS"], "wbs"), *ids(backlog_reverse.get(row["ID"], ""), "wbs")])
        dec_reverse = reverse(wbs, "결정")
        for row in decisions:
            row["WBS"] = joined([*ids(row["WBS"], "wbs"), *ids(dec_reverse.get(row["ID"], ""), "wbs")])
        # Body-explicit backlog links can establish a decision-to-work relationship.
        back_wbs = {r["ID"]: ids(r["WBS"], "wbs") for r in backlog}
        for row in decisions:
            row["WBS"] = joined([*ids(row["WBS"], "wbs"), *(x for b in ids(row["백로그"], "backlog") for x in back_wbs.get(b, []))])
        decision_wbs = {r["ID"]: ids(r["WBS"], "wbs") for r in decisions}
        commits = []
        for record in git("log", "main", "--reverse", "--topo-order", "--format=%H%x1f%aI%x1f%an%x1f%P%x1f%B%x1e").split("\x1e"):
            if not record.strip():
                continue
            sha, when, author, parents, message = record.strip().split("\x1f", 4)
            subject = message.splitlines()[0]
            conventional = re.match(r"([a-z]+)(?:\(([^)]+)\))?!?:\s*(.*)", subject)
            ds = self.decision_ids(message)
            commits.append({"커밋": sha[:7], "날짜": when[:16].replace("T", " "), "작성자": author, "유형": conventional[1] if conventional else "",
                            "범위": conventional[2] or "" if conventional else "", "요약": conventional[3] if conventional else subject,
                            "결정": joined(ds), "백로그": joined(ids(message, "backlog")),
                            "WBS": joined([*ids(message, "wbs"), *(x for d in ds for x in decision_wbs.get(d, []))]),
                            "병합": "병합" if len(parents.split()) > 1 else "", "기기 배포": "",
                            "링크": f"https://github.com/MCP-governance/mcp-gateway/commit/{sha}"})
        commits.sort(key=lambda r: datetime.fromisoformat(r["날짜"]))
        self.add("commits", commits, ["git log main"], links={"결정": "decisions", "백로그": "backlog", "WBS": "wbs", "기기 배포": "deploys"})
        dec_commits = defaultdict(list)
        for row in commits:
            for d in row["결정"].split(SEP):
                if d:
                    dec_commits[d].append(row["커밋"])
        for row in decisions:
            row["커밋"] = joined(dec_commits[row["ID"]])

    def originals_and_policies(self):
        mvp_path = next((WIKI / "analysis/development-wbs/originals").glob("*MVP_WBS_2026-10-03.xlsx"))
        headers, mvp = workbook_table(mvp_path, "MVP WBS")
        curation_path = WIKI / "analysis/drive-curation/README.md"
        mvp_links = {int(r["원본 행"]): r for _, rows in md_tables(read(curation_path)) for r in rows if "원본 행" in r and r["원본 행"].isdigit()}
        for r in mvp:
            link = mvp_links.get(r["_row"], {})
            r["ID"] = f"MVP-{r['_row'] - 4:02d}"
            r["개발 목표"] = joined(re.findall(r"G\d{2}", link.get("목표", "")))
            r["남은 확인"] = link.get("남은 확인", "")
        self.add("mvp", mvp, [mvp_path, curation_path], ["ID", *headers, "개발 목표", "남은 확인"])
        plan_path = next((WIKI / "threats/originals").glob("*V1.5.xlsm"))
        headers, rows = workbook_table(plan_path, "WBS", plan=True)
        for row in rows:
            row["ID"] = f"PLAN-{row['_row']:03d}"
        self.add("plan-v15", rows, [plan_path], ["ID", *headers])
        policy_path = next((WIKI / "threats/originals").glob("*50_2026-10-14.xlsm"))
        for slug, sheet in [("policies-detail", "정책 최종본"), ("policies-docs", "정책 최종본(독스용)")]:
            headers, rows = workbook_table(policy_path, sheet)
            self.add(slug, rows, [policy_path], headers)
        extra_path = next((WIKI / "threats/originals").glob("*25.xlsm"))
        headers, rows = workbook_table(extra_path, "(본)정책 최종본")
        self.add("policies-extra", rows, [extra_path], headers)
        _, policies = workbook_table(policy_path, "(본)정책 최종본")
        spec_path = WIKI / "analysis/feature-spec-2026-10-17/feature-spec.json"
        coverage = {r["policy"]: r for r in load(spec_path)["policy_coverage"]}
        basis_path = WIKI / "analysis/compliance/pol-legal-basis.json"
        basis = {r["policy_id"]: r for r in load(basis_path)["policies"]}
        mapping_path = WIKI / "analysis/drive-curation/policy-id-mapping.json"
        mappings = load(mapping_path)["mappings"]
        by_id = {r["policy_id"]: r for r in mappings}
        baseline_path = WIKI / "analysis/development-wbs/snapshot-before-wbs-1017.json"
        baseline_ids = {r["id"] for r in load(baseline_path)["items"]}
        cross_path = WIKI / "analysis/drive-curation/policy-50-crosswalk.json"
        crosses = {r["policy_id"]: r for r in load(cross_path)}
        feature_reverse = reverse(self.tables["features"]["rows"], "정책")
        wbs_reverse = reverse(self.tables["wbs"]["rows"], "정책")
        result = []
        for r in policies:
            ident = r["ID"]
            c, b, mapping = coverage.get(ident, {}), basis.get(ident, {}), by_id.get(ident, {})
            status = mapping.get("status_report", {})
            evidence = "\n".join(f"{k}: {v}" for k, v in r.items() if k not in {"ID", "영역", "정책명", "정책 내용", "CIA 영향", "_row"} and v)
            evidence += "\n" + url(repo_path(cross_path)) + ": " + dump(crosses.get(ident, {}))
            result.append({"ID": ident, "영역": r["영역"], "정책명": r["정책명"], "정책 내용": r["정책 내용"], "CIA 영향": r["CIA 영향"],
                           "자동화 분류": c.get("class", ""), "현재": c.get("now", status.get("current", "")),
                           "개발 후": c.get("after", status.get("after_planned_development", "")),
                           "관련 추가 개발": joined(ids(c.get("planned", "") or status.get("planned_wbs_ref", ""), "wbs")),
                           "자동화 한계": c.get("limit", status.get("automation_limit", "")), "기능": feature_reverse.get(ident, ""),
                           "WBS": wbs_reverse.get(ident, ""), "규제 조항 수": b.get("provisions", ""), "구속력 의무": b.get("binding_mandatory", ""),
                           "요구 ID": joined(b.get("requirements", [])), "번호 대응": ident, "근거": evidence})
        self.add("policies", result, [policy_path, spec_path, basis_path, mapping_path, cross_path],
                 links={"기능": "features", "WBS": "wbs", "관련 추가 개발": "wbs", "요구 ID": "reg-requirements", "번호 대응": "policy-id-map"})
        rows = []
        for r in mappings:
            row = {"ID": r["policy_id"], "정책명": r["name"], "정책": r["policy_id"], "세부": r}
            for system, label in [("pac50-docs", "독스용 번호"), ("pac24-original", "원본 번호"), ("runtime-pac15", "실행 번호")]:
                row[label] = joined(x["id"] for x in r["links"] if x["system"] == system)
            row["검수 작업"] = joined(x for x in ids(wbs_reverse.get(r["policy_id"], ""), "wbs") if x in baseline_ids)
            rows.append(row)
        self.add("policy-id-map", rows, [mapping_path, spec_path, baseline_path], ["ID", "정책명", "정책", "독스용 번호", "원본 번호", "실행 번호", "검수 작업", "세부"],
                 {"정책": "policies", "독스용 번호": "policies-docs", "원본 번호": "policies-detail", "검수 작업": "wbs"})

    def experiments(self):
        source = WIKI / "catalogs/experiments.json"
        rows, sources = [], [source]
        task_reverse = reverse(self.tables["wbs"]["rows"], "시험")
        for r in load(source):
            card = WIKI / r["path"]
            text = read(card)
            sources.append(card)
            heading = re.search(r"^#\s+(.*)$", text, re.M)
            title = re.sub(r"^E\d+\s*[—:·-]\s*", "", heading[1]) if heading else r["title"]
            explicit = {key: ids([r, text], key) for key in ["policies", "wbs", "backlog"]}
            rows.append({"ID": r["id"], "이름": title, "목적": r.get("summary", ""), "조건": r.get("prerequisites", []),
                         "판정 기준": r.get("expected", {}), "상태": VALUES.get(r.get("status"), r.get("status", "")),
                         "정책": joined(explicit["policies"]), "WBS": joined([*explicit["wbs"], *ids(task_reverse.get(r["id"], ""), "wbs")]),
                         "백로그": joined(explicit["backlog"]), "카드": url(repo_path(card))})
            status_line = next((line for line in text.splitlines()[:12] if re.search(r"\*\*(?:미실행|진행 중|완료|보류)\*\*", line)), "")
            if status_line:
                rows[-1]["상태"] = re.search(r"\*\*(미실행|진행 중|완료|보류)\*\*", status_line)[1]
        self.add("experiments", rows, sources, links={"정책": "policies", "WBS": "wbs", "백로그": "backlog"})

    def deployments(self):
        sources = [WIKI / "analysis/wbs-1015/deploy.json", *sorted((WIKI / "analysis/wbs-1017").glob("*")),
                   WIKI / "status/current.md", *sorted((WIKI / "status/reports").glob("*.md"))]
        records = {}

        def add(when, target, sha, content, source, branch=""):
            if not when or not re.fullmatch(r"[0-9a-f]{7,40}", sha or ""):
                return
            key = f"{when[:10]}|{target}|{sha[:7]}"
            if key not in records:
                records[key] = {"날짜": when[:10], "대상": target, "커밋": sha[:7], "브랜치": branch,
                                "내용": content, "근거 자료": []}
            row = records[key]
            row["근거 자료"].append(url(repo_path(source)))
            if branch and not row["브랜치"]:
                row["브랜치"] = branch

        for source in sources:
            if source.suffix == ".json":
                r = load(source)
                when = r.get("as_of", r.get("audited_at", ""))
                if source.name == "deploy.json" and r.get("field_up_rc") == 0:
                    add(when, "솔루션 기기", r["commit"], r.get("subject", "field up"), source)
                dep = r.get("deployment", {})
                if dep and (dep.get("rc") == 0 or any("field up rc 0" in x for x in dep.get("steps", []))):
                    add(when, "솔루션 기기", dep.get("commit"), dump(dep), source)
                if r.get("install", {}).get("result", "").startswith("previous managed-pj1 removed and reinstalled"):
                    add(when, "PJ1", r.get("kit", {}).get("server_commit"), r["install"]["result"], source)
                if r.get("pj1", {}).get("endpoint") and "active" in r["pj1"]["endpoint"]:
                    add(when, "PJ1", dep.get("commit"), dump(r["pj1"]), source)
                continue
            if source.suffix != ".md":
                continue
            text = read(source)
            date_match = re.search(r"20\d{2}-\d{2}-\d{2}", source.name)
            when = date_match[0] if date_match else ""
            title_date = re.search(r"^# .*?(20\d{2}-\d{2}-\d{2})", text, re.M)
            if title_date:
                when = title_date[1]
            current_sha = ""
            source_is_field = bool(re.search(r"솔루션 기기|기기에서|field up|운영자가 띄웠던", text))
            for line in text.splitlines():
                if source.name == "current.md" and line.startswith("## "):
                    # Each current-status section can describe a different historical date.
                    when = ""
                if line.startswith("#"):
                    shas = re.findall(r"(?<![0-9a-f])[0-9a-f]{7}(?![0-9a-f])", line)
                    current_sha = shas[0] if len(shas) == 1 else ""
                if not re.search(r"배포|설치|재설치", line):
                    continue
                branch = re.search(r"(?:통합 브랜치|브랜치)\s*`((?:\d{2}/\d{2}-|feat/)[^`\s]+)`", line)
                branch_name = branch[1] if branch else ""
                sha_pattern = r"`?([0-9a-f]{7})`?(?![0-9a-f])"
                labeled = re.findall(r"(?:최종 배포|\d차 배포)\s*" + sha_pattern, line)
                phased = re.search(r"솔루션 기기 배포\s+(.+)", line)
                if phased:
                    labeled += re.findall(r"\d차\s*" + sha_pattern, phased[1])
                interim = re.search(r"중간 배포\s+((?:`[0-9a-f]{7}`[·\s]*)+)", line)
                if interim:
                    labeled += re.findall(r"[0-9a-f]{7}", interim[1])
                labeled += re.findall(sha_pattern + r"\(배포\)", line)
                if labeled and source_is_field:
                    for sha in unique(labeled):
                        add(when, "솔루션 기기", sha, markdown_links(line.strip("- |"), source), source, branch_name)
                    continue
                successful_install = "설치 완료" in line and "실패" not in line
                if re.search(r"미배포|배포[^.。;]*?(?:미실행|하지 않|하지 않았|없음|금지|예정|대기)|설치[^.。;]*?(?:미실행|미완료|실패|필요|대기)|배포 전|배포 명령", line):
                    if not successful_install:
                        continue
                affirmative = re.search(r"배포(?:[했한됨·: ]|\()|설치 완료|재설치.*(?:완료|active|활성화)|승인 배포.*설치", line)
                if not affirmative:
                    continue
                target = "pj3" if re.search(r"pj3", line, re.I) or "pj3-windows" in source.name else (
                    "PJ1" if re.search(r"PJ1.*(?:설치|재설치)|(?:설치|재설치).*PJ1", line, re.I) else (
                        "솔루션 기기" if re.search(r"솔루션|기기|field up", line) else ""))
                if not target:
                    continue
                shas = unique(re.findall(r"(?<![0-9a-f])[0-9a-f]{7}(?![0-9a-f])", line))
                if not shas and current_sha:
                    shas = [current_sha]
                # Multiple hashes on a line can be merge members rather than deployed revisions.
                if len(shas) != 1:
                    continue
                add(when, target, shas[0], markdown_links(line.strip("- |"), source), source, branch_name)
        mapping = self.stable_ids("deploys", records, "DEP-", 2)
        rows = []
        for key, row in sorted(records.items()):
            row["ID"] = mapping[key]
            row["근거 자료"] = joined(row["근거 자료"])
            rows.append(row)
        self.add("deploys", rows, sources, links={"커밋": "commits"})
        by_sha = defaultdict(list)
        for row in rows:
            by_sha[row["커밋"]].append(row["ID"])
        for row in self.tables["commits"]["rows"]:
            row["기기 배포"] = joined(by_sha[row["커밋"]])

    def compliance(self):
        source = WIKI / "analysis/compliance/requirements.json"
        plan_source = WIKI / "analysis/compliance/dev-plan/dev-plan.json"
        book_path = WIKI / "analysis/compliance/compliance-requirements.xlsx"
        plan_book = WIKI / "analysis/compliance/dev-plan/dev-plan.xlsx"
        if source.exists():
            req = load(source)
            docs, requirements = req["documents"], req["requirements"]
            self.generic("reg-docs", docs, source, ["title", "jurisdiction", "instrument_type", "binding", "status", "key_dates", "provisions", "official_page_url", "page"])
            self.generic("reg-requirements", requirements, source,
                         ["doc_id", "jurisdiction", "ref", "topic_ko", "strength", "binding_doc", "actor", "timing_ko", "obligation_ko", "gateway_control_ko", "pol_ids", "binding_mandatory"],
                         id_key="req_id", links={"규제 문서": "reg-docs", "정책": "policies"})
            headers, rows = workbook_table(book_path, "주제×관할")
            for row in rows:
                row["ID"] = row[headers[0]]
            self.add("reg-matrix", rows, [book_path], ["ID", *headers])
        else:
            for slug, sheet in [("reg-docs", "문서"), ("reg-requirements", "요구사항"), ("reg-matrix", "주제×관할")]:
                headers, rows = workbook_table(book_path, sheet)
                for row in rows:
                    row["ID"] = row[headers[0]]
                self.add(slug, rows, [book_path], ["ID", *[h for h in headers if h != "ID"]])
            docs, requirements = [], []
        if not plan_source.exists():
            names = ["요약", "능력", "개발 제안", "기존 항목 보강", "결정", "첫 스프린트", "사고 기한표", "보존 기간표", "관할 요약", "근거 조항"]
            for slug, sheet in zip([t[2] for t in TABS if 23 <= t[0] <= 32], names):
                if sheet == "요약":
                    book = load_workbook(plan_book, read_only=True, data_only=True, keep_vba=False)
                    try:
                        vals = list(book[sheet].values)
                    finally:
                        book.close()
                    headers, rows = list(vals[0]), [dict(zip(vals[0], r)) for r in vals[1:] if any(v is not None for v in r)]
                else:
                    headers, rows = workbook_table(plan_book, sheet)
                for row in rows:
                    row["ID"] = joined([row[headers[0]], row.get("요구 ID", "")]) if sheet == "근거 조항" else row[headers[0]]
                self.add(slug, rows, [plan_book], ["ID", *[h for h in headers if h != "ID"]])
            return
        p = load(plan_source)
        summary = [("작성일", p["generated_on"]), ("근거 문서", len(docs)), ("요구사항(조항)", len(requirements)),
                   ("구속력 의무 조항", sum(bool(r["binding_mandatory"]) for r in requirements)),
                   ("추가 개발 제안", len(p["proposals"])), ("결정 필요", len(p["decisions"])), ("첫 스프린트 작업", len(p["sprint1"]))]
        self.add("reg-plan-summary", [{"ID": k, "값": v} for k, v in summary], [plan_source, source], ["ID", "값"])
        self.generic("reg-capabilities", p["capabilities"], plan_source, ["title", "demand", "existing", "gap", "proposal"], links={"제안": "reg-proposals"})
        rows = self.generic("reg-proposals", p["proposals"], plan_source, ["title", "capability", "extends", "scope", "acceptance", "depends", "week", "decisions", "risks"],
                            links={"백로그": "backlog", "능력": "reg-capabilities", "선행": "backlog", "결정 대기": "reg-decisions"})
        for row in rows:
            row["백로그"] = self.registry.get("backlog_aliases", {}).get(row["ID"], joined(ids(row["ID"], "backlog")))
        self.tables["reg-proposals"]["headers"].append("백로그")
        self.generic("reg-reinforcements", p["acceptance_additions"], plan_source, ["target", "add"], id_key="dev", links={"ID": "backlog"})
        self.generic("reg-decisions", p["decisions"], plan_source, ["question", "options", "recommendation", "needed_by"])
        rows = self.generic("reg-sprint", p["sprint1"], plan_source, ["task", "done_when", "dev"], id_key="n", links={"백로그": "backlog"})
        for row in rows:
            row["백로그"] = explicit_list(row["백로그"], "DEV-")
        req_map = {r["req_id"]: r for r in requirements}
        for slug, key in [("reg-incident-deadlines", "deadlines"), ("reg-retention", "retention")]:
            rows = [{**r, "timing": r.get("timing", req_map.get(r["req_id"], {}).get("timing_ko", ""))} for r in p[key]]
            self.generic(slug, rows, plan_source, ["doc_id", "jurisdiction", "ref", "timing", "topic", "strength", "binding_doc", "doc_status", "pol_ids"],
                         id_key="req_id", links={"ID": "reg-requirements", "규제 문서": "reg-docs", "정책": "policies"})
        rows = []
        for group in sorted({d["jurisdiction_group"] for d in docs}, key=lambda g: (g != "KR", g)):
            ds = [d for d in docs if d["jurisdiction_group"] == group]
            rs = [r for r in requirements if r["jurisdiction_group"] == group]
            rows.append({"ID": group, "관할": JUR.get(group, group), "문서 수": len(ds), "구속력 문서": sum(bool(d["binding"]) for d in ds),
                         "조항 수": len(rs), "구속력 의무": sum(bool(r["binding_mandatory"]) for r in rs),
                         "시행 전·초안": joined(d["id"] for d in ds if d["status"] in {"draft", "adopted_not_yet_applicable", "proposed"})})
        self.add("reg-jurisdictions", rows, [source], ["ID", "관할", "문서 수", "구속력 문서", "조항 수", "구속력 의무", "시행 전·초안"], {"시행 전·초안": "reg-docs"})
        rows = []
        for cap in p["capabilities"]:
            for r in cap["evidence_rows"]:
                rows.append({"ID": f"{cap['id']}:{r['req_id']}", "능력": cap["id"], "요구 ID": r["req_id"], "규제 문서": r["doc_id"],
                             "관할": JUR.get(r["jurisdiction"], r["jurisdiction"]), "조항": r["ref"], "요지": r["topic"],
                             "강도": VALUES.get(r["strength"], r["strength"]), "구속력": r["binding_doc"], "문서 상태": VALUES.get(r["doc_status"], r["doc_status"])})
        self.add("reg-clauses", rows, [plan_source], ["ID", "능력", "요구 ID", "규제 문서", "관할", "조항", "요지", "강도", "구속력", "문서 상태"],
                 {"능력": "reg-capabilities", "요구 ID": "reg-requirements", "규제 문서": "reg-docs"})

    def references(self):
        configs = [
            ("ent-requirements", "analysis/requirements.json", "area requirement acceptance negative_cases status priority gap code_evidence test_evidence source_refs", {}),
            ("competitor-map", "analysis/competitor-mapping.json", "requirement_id product feature version_or_commit verification our_gap evidence_url evidence_paths analysis_page", {"요구 ID": "ent-requirements"}),
            ("integrations", "analysis/integrations.json", "product category version_or_commit route auth capabilities license compatibility_status gap source_refs", {}),
            ("integration-cases", "analysis/integration-cases.json", "product feature prerequisites procedure expected negative_case status evidence failure_fix code_evidence source_refs", {}),
            ("mcp-candidates", "analysis/mcp-candidates.json", "title package version transport auth_boundary license experiment_ids execution_status verification source_url limitations", {"시험": "experiments"}),
            ("skill-candidates", "analysis/skill-candidates.json", "title repository commit license status reuse_plan unreviewed_scope source_url verification", {}),
            ("oss-reuse", "analysis/oss-reuse.json", "title repository commit license existing_source_path status implementation_guidance selection_gate experiment_ids source_url verification", {"시험": "experiments"}),
            ("solutions", "catalogs/solutions.json", "title kind commit version license summary verification verified_on url path", {}),
            ("reference-updates", "analysis/reference-updates.json", "title doi publication_status venue published_on reading_scope methodology_adaptation adoption_status experiment_ids existing_catalog_id source_url verification", {"시험": "experiments", "기존 자료": "papers"}),
            ("imports", "catalogs/imports.json", "source_project upstream_id imported_on files distribution provenance", {"원천 ID": "papers"}),
            ("threats", "catalogs/threats.json", "title source_id threat boundary runtime_candidates code_candidates verification limitations local_tool_risk_ids", {"원본 ID": "policies-detail", "위험 ID": "local-tool-risks"}),
            ("local-tool-risks", "catalogs/local-tool-risks.json", "title source_controls tools boundary observation verification required_control remaining path", {"원본 통제": "policies-detail"}),
            ("endpoint-permission-refs", "catalogs/endpoint-permission-references.json", "title url version commit license obtained read_scope official_claim application validation_needed limitations", {}),
        ]
        for slug, name, fields, links in configs:
            source = WIKI / name
            self.generic(slug, load(source), source, fields.split(), links=links)
        deep = sorted((WIKI / "analysis/oss-deep").glob("*.md"))
        # Only literal repository/product mentions connect a deep analysis document.
        for row in self.tables["oss-reuse"]["rows"]:
            repository = row["저장소"]
            row["심층 분석"] = linked_paths(repo_path(p) for p in deep if repository.lower() in read(p).lower())
        self.tables["oss-reuse"]["headers"].append("심층 분석")
        self.tables["oss-reuse"]["sources"].extend(repo_path(p) for p in deep)
        source = WIKI / "catalogs/papers.json"
        self.generic("papers", load(source), source,
                     "title authors year publication_status venue doi url engineering_use reading_scope verification verified_on license_url page".split(), details=False)
        source = WIKI / "analysis/code-structure.json"
        structure = load(source)
        basis = structure["metadata"]["head"]
        rows = []
        qualified_sources = []
        for module in structure["modules"]:
            counts = Counter(f["symbol"] for f in module.get("functions", []))
            qualified = {}
            if any(n > 1 for n in counts.values()):
                # AST supplies class/nesting names missing from the saved function list.
                def visit(node, parents=()):
                    for child in ast.iter_child_nodes(node):
                        scoped = isinstance(child, (ast.ClassDef, ast.FunctionDef, ast.AsyncFunctionDef))
                        names = (*parents, child.name) if scoped else parents
                        if isinstance(child, (ast.FunctionDef, ast.AsyncFunctionDef)):
                            qualified[(child.name, child.lineno)] = ".".join(names)
                        visit(child, names)
                visit(ast.parse(git("show", f"{basis}:{module['path']}")))
                qualified_sources.append(f"git show {basis}:{module['path']}")
            for function in module.get("functions", []):
                symbol = function["symbol"]
                if counts[symbol] > 1:
                    symbol = qualified.get((symbol, function["line_start"]))
                    if symbol is None:
                        raise ValueError(f"함수 이름 범위 대조 실패: {module['path']}:{function['line_start']}")
                rows.append({"ID": f"{module['path']}::{symbol}", "파일": url(module["path"]), "함수": symbol,
                             "시작 행": function["line_start"], "끝 행": function["line_end"], "비동기": function.get("async", False),
                             "호출": function.get("call_expressions", []), "소스 상태": module.get("source_state", ""),
                             "해시": module.get("sha256", ""), "기준 커밋": basis[:7], "세부": {k: v for k, v in function.items() if k not in {"symbol", "line_start", "line_end", "async", "call_expressions"}}})
        self.add("code-structure", rows, [source, *qualified_sources], ["ID", "파일", "함수", "시작 행", "끝 행", "비동기", "호출", "소스 상태", "해시", "기준 커밋", "세부"], {"기준 커밋": "commits"})
        enterprise = next((WIKI / "analysis/outputs/01a13365-enterprise").glob("*.xlsx"))
        book = load_workbook(enterprise, read_only=True, data_only=True, keep_vba=False)
        try:
            names = book.sheetnames
        finally:
            book.close()
        absorbed = {"요구사항": "ent-requirements", "코드구조": "code-structure", "기능대조": "competitor-map", "연동대상": "integrations",
                    "연동시험": "integration-cases", "MCP대상": "mcp-candidates", "실험계획": "experiments", "연구근거": "reference-updates", "OSS스킬": "skill-candidates / oss-reuse"}
        self.notes.append("기업 검증대장 흡수: " + "; ".join(f"{a} → {b}" for a, b in absorbed.items()) + ".")
        for sheet in names:
            if sheet in absorbed:
                continue
            slug = "enterprise-summary" if sheet == "현황" else "enterprise-" + hashlib.sha256(sheet.encode()).hexdigest()[:8]
            headers, rows = workbook_table(enterprise, sheet)
            for row in rows:
                row["ID"] = str(row.get(headers[0]) or row.get(headers[3]) or f"행-{row['_row']}")
            headers = [h for h in headers if any(r.get(h) not in (None, "") for r in rows)]
            self.add(slug, rows, [enterprise], ["ID", *headers])
            self.notes.append(f"추가 탭 {slug}: 기업 검증대장의 ‘{sheet}’는 별도 집계 표여서 보존.")

    def documents(self):
        last_dates = {}
        for record in git("log", "HEAD", "--format=%x1e%aI", "--name-only", "--", "research/llm-wiki", "docs/ai").split("\x1e"):
            parts = record.strip().splitlines()
            if not parts:
                continue
            for path in parts[1:]:
                if path:
                    last_dates.setdefault(path, parts[0][:10])
        paths = set()
        for path in WIKI.rglob("*"):
            if not path.is_file():
                continue
            rel = path.relative_to(WIKI)
            if "sheet" == rel.parts[0]:
                continue
            if path.suffix.lower() == ".md" and rel.parts[0] in {"sources", "local-sources"} and rel.as_posix() != "sources/README.md":
                continue
            if path.suffix.lower() == ".md":
                paths.add(path)
            if "originals" in rel.parts or path.suffix.lower() == ".xlsx":
                # sources/local-sources originals are excluded only as Markdown copies.
                paths.add(path)
        paths.update((ROOT / "docs/ai").glob("*.md"))
        records = {}
        for path in sorted(paths):
            rel = repo_path(path)
            text = read(path) if path.suffix.lower() == ".md" else ""
            title = re.search(r"^# (.+)$", text, re.M)
            when = filename_date(path.name) or last_dates.get(rel, "")
            kind = "기타"
            if "/originals/" in rel or path.suffix.lower() in {".xlsx", ".xlsm"}:
                kind = "원본"
            elif rel.startswith("docs/ai/"):
                kind = "결정" if path.name == "DECISIONS.md" else "설계 정본"
            else:
                part = path.relative_to(WIKI).parts[0]
                kind = {"status": "보고" if "/reports/" in rel else "현황", "analysis": "분석", "experiments": "실험",
                        "methodology": "방법론", "catalogs": "카탈로그", "papers": "레퍼런스", "solutions": "제품 분석",
                        "threats": "위협·정책", "integrations": "기업 연동", "templates": "양식", "sources": "원본",
                        "tools": "도구"}.get(part, "안내" if len(path.relative_to(WIKI).parts) == 1 else "기타")
            state = "정본"
            if kind in {"레퍼런스", "제품 분석"}:
                state = "참고"
            elif kind == "원본":
                state = "보존"
            elif kind == "보고" or (kind == "분석" and path.name != "README.md") or                     re.search(r"history|snapshot|이력|스냅샷", path.name, re.I):
                state = "기록"
            if replaced_document(text):
                state = "대체됨"
            related = [x for x in ids(text) if re.fullmatch(r"G\d{2}\.\d{2}|DEV-\d{2}|D-\d+|POL-\d{2}|E\d{2}", x)]
            records[rel] = {"제목": title[1].strip() if title else path.name, "종류": kind, "날짜": when,
                            "경로": url(rel), "관련 ID": joined(related[:15]), "상태": state}
        drive_source = WIKI / "team-agi-deliverables-2026-10-15.md"
        drive_text = read(drive_source)
        titles = {m[2]: m[1] for m in re.finditer(r"\[([^\]]+)\]\((https://(?:drive|docs)\.google\.com/[^)]+)\)", drive_text)}
        for link in unique(re.findall(r"https://(?:drive|docs)\.google\.com/[^\s<>)\]|]+", drive_text)):
            records[link] = {"제목": titles.get(link, "Drive 자료"), "종류": "외부", "날짜": "2026-10-15", "경로": link, "관련 ID": "", "상태": "기록"}
        mapping = self.stable_ids("docs", records, "DOC-", 3)
        rows = [{"ID": mapping[key], **record} for key, record in sorted(records.items())]
        self.add("docs", rows, [*sorted(paths), drive_source])

    def build(self):
        self.management()
        self.originals_and_policies()
        self.experiments()
        self.deployments()
        self.compliance()
        self.references()
        self.documents()
        return self


def integrity(manifest, payloads):
    result = {"tabs": {}, "duplicate_ids": 0, "format_errors": 0, "missing_links": 0}
    tables = {}
    for tab in manifest["tabs"]:
        if not tab["csv"]:
            continue
        raw = payloads[tab["csv"]]
        parsed = list(csv.reader(io.StringIO(raw.decode("utf-8"), newline="")))
        headers, values = parsed[0], parsed[1:]
        tables[tab["slug"]] = [dict(zip(headers, row)) for row in values]
    for tab in manifest["tabs"]:
        slug = tab["slug"]
        if not tab["csv"]:
            continue
        raw, rows = payloads[tab["csv"]], tables[slug]
        errors = []
        header = [c["name"] for c in tab["columns"]]
        parsed = list(csv.reader(io.StringIO(raw.decode("utf-8"), newline="")))
        if raw.startswith(b"\xef\xbb\xbf") or b"\r" in raw or parsed[0] != header or len(set(header)) != len(header):
            errors.append("CSV 머리글·인코딩·줄바꿈 오류")
        if len(rows) != tab["rows"]:
            errors.append("행 수 오류")
        if any(len(row) != len(header) for row in parsed[1:]):
            errors.append("열 수 오류")
        counts = Counter(row[tab["id_column"]] for row in rows)
        duplicates = sorted(k for k, v in counts.items() if v > 1)
        missing = []
        for row in rows:
            ident = row[tab["id_column"]]
            if not ident or (slug in PATTERNS and not re.fullmatch(PATTERNS[slug], ident)):
                errors.append(f"ID 형식: {ident}")
            for key, value in row.items():
                if len(value) > MAX_CELL or re.match(r"^\s*[=+@]", value):
                    errors.append(f"셀 형식: {ident}/{key}")
                if key in FORMULA_COLUMNS.get(slug, ()) and value:
                    errors.append(f"수식 열은 빈 값: {ident}/{key}")
            if slug == "wbs":
                for key in ["구현", "회귀", "배포", "실기기", "검수"]:
                    if row[key] not in {"0", "0.0", "0.5", "1", "1.0"}:
                        errors.append(f"점수 형식: {ident}/{key}")
        link_report = {}
        for column, target in tab["links"].items():
            if column not in header or target not in tables:
                errors.append(f"연결 정의 오류: {column} → {target}")
                continue
            target_tab = next(t for t in manifest["tabs"] if t["slug"] == target)
            target_ids = {r[target_tab["id_column"]] for r in tables[target]}
            absent = []
            for row in rows:
                for token in row[column].split(SEP):
                    if not token:
                        continue
                    if target in PATTERNS and not re.fullmatch(PATTERNS[target], token):
                        errors.append(f"연결 ID 형식: {row[tab['id_column']]}/{column}/{token}")
                    if token not in target_ids:
                        absent.append({"row_id": row[tab["id_column"]], "id": token})
            link_report[column] = {"target": target, "missing_ids": sorted({r["id"] for r in absent}), "missing": absent}
            missing.extend(absent)
        result["tabs"][slug] = {"rows": len(rows), "duplicate_ids": duplicates, "format_errors": errors, "links": link_report,
                                "sha256": hashlib.sha256(raw).hexdigest()}
        result["duplicate_ids"] += len(duplicates)
        result["format_errors"] += len(errors)
        result["missing_links"] += len(missing)
    return result


def output(export, flat):
    manifest = {"generated_at": git("show", "-s", "--format=%cI", "HEAD"), "source_head": git("rev-parse", "HEAD"),
                "raw_base": RAW, "tabs": [], "id_registry": export.registry, "flat": flat}
    payloads = {}
    extras = [slug for slug in export.tables if slug not in {t[2] for t in TABS}]
    tab_specs = [*TABS, *[(46 + i, "기업 현황" if slug == "enterprise-summary" else slug, slug, "기업", "참조")
                          for i, slug in enumerate(extras)]]
    for order, title, slug, group, kind in tab_specs:
        table = export.tables.get(slug)
        tab = {"order": order, "title": title, "slug": slug, "group": group, "kind": kind,
               "csv": f"tabs/{order:02d}-{slug}.csv" if table else None,
               "rows": len(table["rows"]) if table else 0,
               "columns": [{"name": h, "formula": h in FORMULA_COLUMNS.get(slug, ())} for h in table["headers"]] if table else [],
               "id_column": "커밋" if slug == "commits" else "ID", "links": table["links"] if table else {}, "sources": table["sources"] if table else []}
        manifest["tabs"].append(tab)
        if table:
            payloads[tab["csv"]] = csv_bytes(table["headers"], table["rows"], flat)
            if slug == "components" and (OUT / "seed/components.csv").exists() and not flat:
                payloads[tab["csv"]] = (OUT / "seed/components.csv").read_bytes().replace(b"\r\n", b"\n")
    report = integrity(manifest, payloads)
    lines = ["# MINT 개발 통합 시트", "",
             "정본은 구글 시트 [MINT 개발 통합 시트](https://docs.google.com/spreadsheets/d/1bfD-wfPn_ldwB-AV2o1RAr6DbAtLnQV6JJQd-exJpLI/edit)이고, "
             "이 폴더는 처음 채우기와 저장소 쪽 스냅샷이다.", "",
             "- 관리 탭(WBS·백로그·결정·배포·자료·기능·구성요소·정책·시험)은 시트에서 고친다. 행은 A열 ID로 찾고, 진척률·상태처럼 수식인 열은 쓰지 않는다.",
             "- 참조·원본 탭은 저장소 원천에서 다시 만든다: 이 스크립트 → 푸시 → 시트 A1의 IMPORTDATA 주소를 새 커밋으로 → 값 고정. 시트에서 고치면 편집 경고가 뜬다.",
             "- 시트 → 저장소: `research/llm-wiki/tools/sheet_pull.py`가 웹 게시된 WBS·백로그 CSV를 이 폴더와 snapshot.json에 반영한다.",
             "- 협업 콘솔 #wbs는 시트의 WBS 값을 5분마다 읽어 snapshot.json 위에 덮어 보여 준다.",
             "- 구조도: [docs/ai/architecture](../../../docs/ai/architecture/README.md)의 노드 ID가 「구성요소」 탭(CMP-)과 같다.", "",
             "| 탭 | 종류 | 행 수 | 원천 | ID |", "| --- | --- | ---: | --- | --- |"]
    for tab in manifest["tabs"]:
        sources = tab["sources"]
        display = [f"[{Path(s).name}]({url(s)})" if not s.startswith("git ") else s for s in sources[:3]]
        if len(sources) > 3:
            display.append(f"외 {len(sources) - 3}개(manifest)")
        lines.append(f"| {tab['order']:02d} {tab['title']} | {tab['kind'] or '—'} | {tab['rows']} | {'; '.join(display) or '—'} | {tab['id_column']} |")
    lines += ["", "ID: 작업 G01.01, 기능 F-01(원본 순서), 정책 POL-01, 백로그 DEV-01/WIKI-01, 결정은 제목의 D-번호(선행 0 포함), 시험 E01, OSS OSS-01, 자료 DOC-001, 배포 DEP-01, 커밋 7자 SHA. 참조표는 원본 ID를 보존한다. 여러 ID는 ` · `로 구분한다.",
              "", "자료·배포 번호는 manifest의 id_registry에 보존하며 새 항목에만 번호를 추가한다. 이 파일과 스냅샷을 함께 보존한다. 원 계획은 원본 행 번호, MVP는 기능 순번, 코드 구조는 파일::함수, 규제 근거는 능력:요구 ID를 쓴다.",
              "", "연구 백로그의 원본 RES-01은 관리 번호 WIKI-02로 옮기고 근거 자료·id_registry에 원본 번호를 보존한다. 규제 개발 제안의 A열은 RES-01을 유지하며 백로그 열로 WIKI-02에 연결한다.",
              "", "연결 열:"]
    for tab in manifest["tabs"]:
        if tab["links"]:
            lines.append(f"- {tab['title']}: " + "; ".join(f"{k} → {v}" for k, v in tab["links"].items()))
    lines += ["", "```powershell", "python -B -X utf8 research/llm-wiki/tools/build_sheet_export.py",
              "python -B -X utf8 research/llm-wiki/tools/build_sheet_export.py --check",
              "python -B -X utf8 research/llm-wiki/tools/build_sheet_export.py --self-test", "```", "",
              "Python 3.12와 openpyxl이 필요하다. `--flat`은 셀 안 줄바꿈을 ` / `로 바꾼다. `--check`는 재생성 없이 원천과 저장된 파일을 대조하고 중복·형식·스냅샷 불일치 시 실패한다. 연결 누락은 원천의 명시 ID를 보존하고 integrity.json에 보고한다.",
              "", "백로그 상태(위에서 먼저 맞는 규칙 적용): 보류·중단·폐기·범위 밖·사용자 결정·결정 대기 → 보류; 미착수·착수 전 → 미착수; 구현·수정·배포·실기기·통과·반영·설치가 있고 미구현·미완료·미실행·미배포·남음·필요·후속·대기·일부가 없으면 구현 완료(검수 전); 초안·설계·조사·검토·계획·정의·문서가 있고 구현·수정·배포·실기기·통과·설치가 없으면 설계·조사; 나머지는 진행 중. 이는 진행 원문 분류이며 단계 점수를 갱신하지 않는다.",
              "", "기능↔WBS: 기능 레코드 또는 리뷰 표의 동일 기능에 작업 ID가 있으면 ‘명시’. 그 외 정책이 겹치면서 기능명이 작업명에 포함되거나, 공통 일반어를 제외한 기능명 토큰 2개(또는 4자 이상 토큰 1개)가 작업명에 있을 때만 ‘정책 일치’. 그 외는 빈 값이다. 연결 근거는 작업 ID마다 기록한다.",
              "", "배포: 성공한 field up, 재설치·활성화 JSON 또는 날짜·대상·SHA·실행 표현이 함께 기록된 문장만 수록한다. 1~3차·중간·최종·(배포)처럼 SHA가 명시된 목록은 각 기록을 보존한다. 계획·미실행·실패는 제외한다. 설치 완료 뒤 활성화 대기는 설치 기록으로 보존한다. 현재 보고서에 기록된 배포와 현재 가동 상태는 별개다.",
              "", "결정 참조의 선행 0은 원문 제목에 맞춘다(D-7 → D-07). 커밋 메시지의 D 번호는 토큰 전체로 대조해 D-12와 D-120을 구분한다. 진척률·상태 수식 열은 빈 값이며 manifest에서 formula=true다.",
              "", "CSV는 UTF-8(BOM 없음), LF, RFC 4180 인용을 사용한다. 선행 =/+/@는 전각 문자로 바꾸며 셀은 45,000자까지 보존한다. 경로는 main 기준 GitHub URL이다. generated_at은 결정적 재생성을 위해 HEAD 커밋 시각이다. raw_base의 <커밋>은 푸시 후 채운다.",
              "", "자료: 원문 사본 Markdown(sources·local-sources)은 제외하되 sources/README.md와 originals의 다른 파일은 포함한다. 생성물 sheet/README.md는 자기 참조를 피하기 위해 자료에서 제외한다. 파일명 날짜를 우선하고 없으면 Git 마지막 수정일을 쓴다. Git에 없는 파일의 날짜는 빈 값이다. 명시적 대체 문구를 우선하며 보고·이력·스냅샷은 기록, 나머지는 정본이다.",
              "", "정책 번호 대응의 검수 작업은 feature-spec의 POL↔WBS 대응 중 기존 40개 스냅샷에 있던 작업만 연결한다. 원본 mapping JSON에는 검수40 직접 연결이 없다. 코드 구조의 동명 메서드는 JSON 기준 커밋의 클래스·중첩 함수명을 붙여 구별한다. 기업 현황은 원본 집계이며 현재 JSON 행 수로 덮어쓰지 않는다.",
              "", *export.notes, ""]
    payloads["README.md"] = "\n".join(lines).encode("utf-8")
    payloads["manifest.json"] = (json.dumps(manifest, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    payloads["integrity.json"] = (json.dumps(report, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    return payloads, report


def self_test():
    assert ids("D-12 D-120 xD-12 D-12x G01.01 DEV-01 E01") == ["D-12", "D-120", "G01.01", "DEV-01", "E01"]
    assert ids("D-120", "decisions") != ["D-12"]
    rows = [{"ID": "F-01", "WBS": "G01.01"}, {"ID": "F-02", "WBS": "G01.01 · G02.01"}]
    assert reverse(rows, "WBS") == {"G01.01": "F-01 · F-02", "G02.01": "F-02"}
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / "small.csv"
        value = '쉼표, "따옴표"\n둘째 줄'
        path.write_bytes(csv_bytes(["ID", "값", "수식", "점수"], [{"ID": "E01", "값": value, "수식": "=1+1", "점수": 0.5}]))
        assert path.read_bytes().count(b'""') == 2
        with path.open(encoding="utf-8", newline="") as handle:
            parsed = list(csv.reader(handle))
        assert parsed[1] == ["E01", value, "＝1+1", "0.5"]
        assert b" / " in csv_bytes(["값"], [{"값": value}], True)
        assert not path.read_bytes().startswith(b"\xef\xbb\xbf")
    assert len(cell("x" * 45_001)) == 45_000 and cell("x" * 45_001).endswith("…")
    assert cell("  +1") == "  ＋1" and cell("@name") == "＠name"
    assert backlog_state("계획, 미착수") == "미착수"
    assert backlog_state("구현·회귀 통과, 실기기 미실행") == "진행 중"
    assert explicit_list("DEV-66·67·69", "DEV-") == "DEV-66 · DEV-67 · DEV-69"
    assert explicit_list("DEV-50~52", "DEV-") == "DEV-50 · DEV-51 · DEV-52"
    assert explicit_list("—", "DEV-") == ""
    assert filename_date("eu-ai-act-2024-1689.pdf") == ""
    assert filename_date("snapshot-2026-10-18.md") == "2026-10-18"
    assert filename_date("original-20261018.json") == "2026-10-18"
    assert not replaced_document('새 결정을 뒤집으면 "→ D-xx로 대체"를 적는다.')
    assert not replaced_document('이 문서가 그 비교다. 다음 후보는 ROADMAP.md로 옮겼다.')
    assert replaced_document('이 문서는 [후속 문서](next.md)로 대체되었다.')
    export = Export.__new__(Export)
    export.registry = {}
    assert export.stable_ids("docs", ["b", "c"], "DOC-", 3) == {"b": "DOC-001", "c": "DOC-002"}
    assert export.stable_ids("docs", ["a", "b", "c"], "DOC-", 3)["b"] == "DOC-001"
    assert "D-120" not in ids("D-12", "decisions")
    print("self-test: PASS (ID·역참조·CSV 인용·수식 방지·셀 상한·상태)")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--flat", action="store_true")
    parser.add_argument("--check", action="store_true")
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()
    if args.self_test:
        self_test()
        return 0
    payloads, report = output(Export().build(), args.flat)
    mismatches = []
    if args.check:
        for name, raw in payloads.items():
            path = OUT / name
            if not path.exists() or path.read_bytes() != raw:
                mismatches.append(name)
        expected = {name for name in payloads if name.startswith("tabs/")}
        mismatches.extend(repo_path(p) for p in (OUT / "tabs").glob("*.csv") if "tabs/" + p.name not in expected)
    else:
        for name, raw in payloads.items():
            path = OUT / name
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(raw)
    manifest = json.loads(payloads["manifest.json"])
    for tab in manifest["tabs"]:
        print(f"{tab['order']:02d} {tab['title']}: {tab['rows']}")
    print(f"{'check' if args.check else 'export'}: duplicates={report['duplicate_ids']}, format_errors={report['format_errors']}, missing_links={report['missing_links']}, mismatches={len(mismatches)}")
    for name in mismatches[:10]:
        print("불일치:", name)
    if report["duplicate_ids"] or report["format_errors"]:
        for slug, details in report["tabs"].items():
            if details["duplicate_ids"] or details["format_errors"]:
                print(slug, details["duplicate_ids"][:5], details["format_errors"][:5])
        return 1
    return 1 if mismatches else 0


if __name__ == "__main__":
    raise SystemExit(main())
