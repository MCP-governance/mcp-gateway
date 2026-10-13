"""Read an XLSM as data, preserve its bytes, and export every populated cell."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def cell(value):
    return text(value).replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def import_book(source, kind):
    source = source.resolve(strict=True)
    target = ROOT / "threats" / "originals" / source.name
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists() and digest(source) != digest(target):
        raise ValueError("Existing original differs: use a new version filename")
    if source != target:
        shutil.copy2(source, target)
    formulas = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    sheets, policies = [], []
    full_md = [f"# {source.name} — 모든 시트 셀 추출", "", "XLSM을 데이터로 읽었습니다. 매크로 실행·수정·재계산을 하지 않았습니다.",
               "수식 문자열과 저장된 캐시 값을 JSON에 별도로 보존합니다. 빈 셀은 생략하고 원래 좌표를 유지합니다.", ""]
    try:
        for sheet in formulas.worksheets:
            cache_rows = dict(enumerate(cached[sheet.title].iter_rows(values_only=True), 1))
            rows, headers = [], None
            full_md += [f"## {sheet.title}", "", "| 행 | 좌표: 내용 |", "| --- | --- |"]
            for row_number, values in enumerate(sheet.iter_rows(), 1):
                nonempty = [c for c in values if c.value is not None]
                if not nonempty:
                    continue
                row = [{"cell": c.coordinate, "value": text(c.value),
                        "formula": c.data_type == "f",
                        "cached_value": text(cache_rows.get(row_number, ())[c.column - 1])
                        if c.column <= len(cache_rows.get(row_number, ())) else ""} for c in nonempty]
                rows.append({"row": row_number, "cells": row})
                full_md.append(f"| {row_number} | " + " · ".join(f"{c['cell']}: {cell(c['value'])}" for c in row) + " |")
                vals = [text(c.value) for c in values]
                if vals and vals[0] in ("정책 번호", "정책 ID"):
                    headers = vals
                elif headers and re.fullmatch(r"(?:PAC|POL-?)\d+", vals[0]):
                    fields = {h.replace("\n", " "): vals[i] if i < len(vals) else ""
                              for i, h in enumerate(headers) if h}
                    policies.append({"source_id": vals[0], "sheet": sheet.title,
                                     "row": row_number, "fields": fields,
                                     "implementation_status": "매핑 검토 대상; 구현 완료 증거 아님"})
            sheets.append({"name": sheet.title, "nonempty_rows": len(rows), "rows": rows})
            full_md.append("")
    finally:
        formulas.close()
        cached.close()
    folder = ROOT / "threats" / "extracted"
    folder.mkdir(parents=True, exist_ok=True)
    stem = "policies" if kind == "policy" else "wbs"
    result = {"source": target.relative_to(ROOT).as_posix(), "sha256": digest(target),
              "imported_on": "2026-10-13", "kind": kind,
              "macros_executed": False, "formulas_recalculated": False,
              "sheets": sheets, "policies": policies}
    (folder / f"{stem}.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (folder / f"{stem}.md").write_text("\n".join(full_md).rstrip() + "\n", encoding="utf-8")
    if kind == "policy":
        for sheet_name, output in (("정책 최종본", "policies-24.md"), ("(본)정책 최종본", "policies-50.md")):
            selected = [p for p in policies if p["sheet"] == sheet_name]
            parts = [f"# {sheet_name} — {len(selected)}개 정책", "",
                     "원본 셀에서 추출한 정책입니다. 이 표의 PAC 번호와 런타임 PAC-01~15는 다른 식별자 체계입니다.",
                     "정책 수와 표준·법령 버전은 원본 기재값이며 현재성이나 구현 완료를 뜻하지 않습니다.", ""]
            for p in selected:
                fields = p["fields"]
                parts += [f"## {p['source_id']} {fields.get('정책명', '')}", "",
                          f"출처: `{sheet_name}!A{p['row']}` · [원본](originals/{source.name})", ""]
                for key, value in fields.items():
                    if value:
                        parts += [f"**{key}**", "", value, ""]
            clean = "\n".join(line.rstrip() for line in "\n".join(parts).splitlines()).rstrip() + "\n"
            (ROOT / "threats" / output).write_text(clean, encoding="utf-8")
    assert digest(source) == digest(target), "Copy integrity failed"
    print(json.dumps({"kind": kind, "sheets": len(sheets), "policy_rows": len(policies),
                      "nonempty_rows": sum(s["nonempty_rows"] for s in sheets), "sha256": result["sha256"]}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--kind", choices=("policy", "wbs"), required=True)
    args = parser.parse_args()
    import_book(args.workbook, args.kind)
