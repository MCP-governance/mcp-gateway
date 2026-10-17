"""Aggregate the compliance notes (analysis/compliance/notes/*.json) into the requirement matrix.

Writes into analysis/compliance/:
  requirements.json            one record per provision, stable keys, with document-level binding/status
  documents.md                 document table
  jurisdiction-theme.md        theme x jurisdiction counts
  by-theme/<theme>.md          requirements per theme, grouped by jurisdiction
  pol-legal-basis.json / .md   which provisions back each of the 50 internal policies
  compliance-requirements.xlsx review workbook (regenerated, not a source of truth)
Run from the repository root:  python -B -X utf8 research/llm-wiki/tools/build_compliance_matrix.py
"""
from __future__ import annotations

import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

WIKI = Path(__file__).resolve().parents[1]
OUT = WIKI / "analysis/compliance"
NOTES = OUT / "notes"
CATALOG = WIKI / "catalogs/papers.json"
POLICIES = WIKI / "threats/extracted/policies.json"
TODAY = "2026-10-17"

THEME_KO = {
    "human_oversight": "인간 감독", "logging_records": "기록·로그", "transparency_labelling": "투명성·표시", "data_protection_lawful_basis": "개인정보 적법 근거",
    "data_minimisation_retention_deletion": "최소화·보유·파기", "cross_border_transfer": "국외 이전", "automated_decision_rights": "자동화된 결정 권리",
    "security_risk_management": "보안 위험 관리", "secure_development_supply_chain": "보안 개발·공급망", "incident_reporting": "사고 보고",
    "access_control_identity": "접근 통제·신원", "third_party_management": "제3자 관리", "risk_impact_assessment": "위험·영향 평가",
    "governance_accountability": "거버넌스·책임", "testing_red_teaming_robustness": "시험·레드티밍·견고성", "agentic_ai": "에이전트 AI",
    "scope_applicability": "적용 범위", "penalties_enforcement": "제재·집행", "other": "기타"}
STRENGTH_KO = {"mandatory": "의무", "recommended": "권고", "informative": "참고"}
JUR_KO = {"KR": "한국", "US-federal": "미국 연방", "US-states": "미국 주", "EU": "EU", "CN": "중국", "CA": "캐나다", "SG": "싱가포르", "UK": "영국", "JP": "일본",
          "AU": "호주", "IN": "인도", "INTL": "국제"}
STATUS_KO = {"in_force": "시행 중", "voluntary": "자율", "draft": "초안", "adopted_not_yet_applicable": "채택·미시행", "proposed": "제안", "superseded": "대체됨",
             "withdrawn": "철회"}


def jur_label(code: str) -> str:
    if code.startswith("US-") and code not in JUR_KO:
        return "미국 주(" + code[3:] + ")"
    return JUR_KO.get(code, code)


def esc(text) -> str:
    return str(text).replace("|", "\\|").replace("\n", " ")


def load_notes():
    notes = []
    for p in sorted(NOTES.glob("*.json")):
        notes.append(json.loads(p.read_text(encoding="utf-8")))
    return notes


def load_reviews():
    path = OUT / "review-log.json"
    return json.loads(path.read_text(encoding="utf-8"))["reviews"] if path.exists() else {}


def review_label(rv):
    if not rv:
        return "미검증"
    d = rv.get("defects", {})
    return f"검증 완료(결함 {d.get('total', 0)}건: 상 {d.get('high', 0)}·중 {d.get('medium', 0)}·하 {d.get('low', 0)})"


def load_policies():
    d = json.loads(POLICIES.read_text(encoding="utf-8"))
    pol = {}
    area = ""
    for r in sorted((r for r in d["policies"] if r.get("sheet") == "(본)정책 최종본"), key=lambda r: r["row"]):
        if not str(r.get("source_id", "")).startswith("POL-"):
            continue
        f = r["fields"]
        area = f.get("영역") or area  # the workbook merges the area cell over the rows of one life-cycle stage
        pol[r["source_id"]] = {"area": area, "title": f.get("정책명", ""), "cia": f.get("CIA 영향", "")}
    return pol


def jurisdiction_group(code: str) -> str:
    # US state instruments carry jurisdiction like "US-CA"; group them under one label for the matrix
    if code.startswith("US-") and code != "US-federal":
        return "US-states"
    return code


def build():
    notes = load_notes()
    catalog = {r["id"]: r for r in json.loads(CATALOG.read_text(encoding="utf-8"))}
    pol = load_policies()
    reqs = []
    docs = []
    reviews = load_reviews()
    for n in notes:
        did = n["id"]
        rec = catalog.get(did, {})
        jg = jurisdiction_group(n["jurisdiction"])
        docs.append({
            "id": did, "title": n["title"], "title_native": n.get("title_native", ""), "issuer": n["issuer"], "jurisdiction": n["jurisdiction"], "jurisdiction_group": jg,
            "instrument_type": n["instrument_type"], "binding": n["binding"], "status": n["status"], "edition": n.get("edition", ""), "key_dates": n["key_dates"],
            "language": n["language"], "distribution": n["distribution"], "official_page_url": n["official_page_url"], "provisions": len(n["provisions"]),
            "page": f"papers/{did}.md", "sha256": (rec.get("files") or [{}])[0].get("sha256", ""), "review": reviews.get(did),
        })
        for i, p in enumerate(n["provisions"], 1):
            key = hashlib.sha1(f"{did}|{p['ref']}|{p['anchor']}".encode("utf-8")).hexdigest()[:8]
            reqs.append({
                "req_id": f"REQ-{did.split('-')[1]}-{i:02d}", "req_key": key, "doc_id": did, "jurisdiction": n["jurisdiction"], "jurisdiction_group": jg, "binding_doc": n["binding"],
                "doc_status": n["status"], "ref": p["ref"], "topic_ko": p["topic_ko"], "obligation_ko": p["obligation_ko"], "actor": p["actor"], "strength": p["strength"],
                "theme": p["theme"], "timing_ko": p.get("timing_ko", ""), "anchor": p["anchor"], "gateway_control_ko": p["gateway_control_ko"], "pol_ids": p["pol_ids"],
                "binding_mandatory": bool(n["binding"] and p["strength"] == "mandatory"),
            })
    return notes, docs, reqs, pol


def write_json(path: Path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")


def write_text(path: Path, text: str):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8", newline="\n")


def doc_link(did: str, rel: str = "../../papers/") -> str:
    return f"[{did}]({rel}{did}.md)"


def md_documents(docs):
    L_ = ["# 컴플라이언스 근거 문서 목록", "", "관할 순서로 정렬한 근거 문서다. 구속력은 문서 전체의 성격이고 조항별 강도는 [요구사항 매트릭스](requirements.json)에 따로 있다.", "",
          "| ID | 문서 | 관할 | 유형 | 구속력 | 상태 | 주요 일자 | 배포 | 조항 수 | 독립 검증 |", "| --- | --- | --- | --- | --- | --- | --- | --- | ---: | --- |"]
    for d in sorted(docs, key=lambda x: (x["jurisdiction_group"], x["id"])):
        L_.append(f"| {doc_link(d['id'])} | {esc(d['title'])} | {jur_label(d['jurisdiction'])} | {d['instrument_type']} | {'있음' if d['binding'] else '없음'} | {STATUS_KO.get(d['status'], d['status'])} | {esc(d['key_dates'])} | {d['distribution']} | {d['provisions']} | {review_label(d['review'])} |")
    return "\n".join(L_) + "\n"


def md_matrix(reqs):
    juris = sorted({r["jurisdiction_group"] for r in reqs}, key=lambda j: (j != "KR", j))
    cnt = defaultdict(Counter)
    bm = defaultdict(Counter)
    for r in reqs:
        cnt[r["theme"]][r["jurisdiction_group"]] += 1
        if r["binding_mandatory"]:
            bm[r["theme"]][r["jurisdiction_group"]] += 1
    L_ = ["# 주제 × 관할 조항 수", "", "칸의 값은 `전체(구속력 있는 문서의 의무 조항)`이다. 빈 칸(—)은 읽은 범위 안에서 해당 조항을 고르지 않았다는 뜻이며 그 관할에 규정이 없다는 뜻이 아니다.", "",
          "| 주제 | " + " | ".join(jur_label(j) for j in juris) + " | 합계 |", "| --- | " + " | ".join("---:" for _ in juris) + " | ---: |"]
    for th in THEME_KO:
        row = []
        total = 0
        for j in juris:
            c, b = cnt[th][j], bm[th][j]
            total += c
            row.append(f"{c}({b})" if c else "—")
        if total:
            L_.append(f"| [{THEME_KO[th]}](by-theme/{th}.md) | " + " | ".join(row) + f" | {total} |")
    return "\n".join(L_) + "\n"


def md_theme(theme, reqs):
    L_ = [f"# {THEME_KO[theme]} ({theme})", "", f"관할 순서로 정리한 요구사항이다. 강도는 문서 안의 표현이고, 문서가 구속력이 없으면 '의무'도 지침 안의 표현이다(구속력은 [문서 목록](../documents.md)). 이 표는 생성물이며 정본은 [requirements.json](../requirements.json)과 각 문서 노트(`notes/`)다.", ""]
    by = defaultdict(list)
    for r in reqs:
        by[r["jurisdiction_group"]].append(r)
    for j in sorted(by, key=lambda j: (j != "KR", j)):
        L_ += [f"## {jur_label(j)}", "", "| 요구 ID | 문서 · 조항 | 요구 내용 | 강도 | 구속력 | 기한 | MINT 통제 | 정책 |", "| --- | --- | --- | --- | --- | --- | --- | --- |"]
        for r in by[j]:
            L_.append(f"| {r['req_id']} | {doc_link(r['doc_id'], '../../../papers/')} {esc(r['ref'])} | **{esc(r['topic_ko'])}** — {esc(r['obligation_ko'])} | {STRENGTH_KO.get(r['strength'], r['strength'])} | {'있음' if r['binding_doc'] else '없음'} | {esc(r['timing_ko']) or '—'} | {esc(r['gateway_control_ko'])} | {esc(', '.join(r['pol_ids'])) or '—'} |")
        L_.append("")
    return "\n".join(L_) + "\n"


def pol_basis(reqs, pol):
    by = defaultdict(list)
    for r in reqs:
        for pid in r["pol_ids"]:
            by[pid].append(r)
    out = []
    for pid in sorted(pol):
        rs = by.get(pid, [])
        rs_sorted = sorted(rs, key=lambda r: (not r["binding_mandatory"], r["strength"] != "mandatory", r["doc_id"]))
        out.append({
            "policy_id": pid, "area": pol[pid]["area"], "title": pol[pid]["title"], "provisions": len(rs), "binding_mandatory": sum(1 for r in rs if r["binding_mandatory"]),
            "jurisdictions": sorted({r["jurisdiction_group"] for r in rs}), "documents": sorted({r["doc_id"] for r in rs}),
            "requirements": [r["req_id"] for r in rs_sorted]})
    return out


def md_pol(basis, req_by_id):
    L_ = ["# 정책 50개별 법적·표준 근거", "", "각 정책(POL-01~50)에 연결된 조항이다. 연결은 노트 작성자가 정책 문구와 조항을 대조해 붙인 것이며 법적 적합성 판단이 아니다. 근거가 없는 정책은 읽은 문서 범위 안에서 연결된 조항이 없다는 뜻이다.", "",
          "| 정책 | 영역 | 정책명 | 연결 조항 | 구속력 문서의 의무 | 관할 | 대표 근거(최대 3) |", "| --- | --- | --- | ---: | ---: | --- | --- |"]
    for b in basis:
        reps = []
        for rid in b["requirements"][:3]:
            r = req_by_id[rid]
            reps.append(f"{r['doc_id']} {esc(r['ref'])}")
        L_.append(f"| {b['policy_id']} | {esc(b['area'] or '—')} | {esc(b['title'])} | {b['provisions']} | {b['binding_mandatory']} | {esc(', '.join(jur_label(j) for j in b['jurisdictions'])) or '—'} | {esc('; '.join(reps)) or '—'} |")
    none = [b["policy_id"] for b in basis if not b["provisions"]]
    L_ += ["", f"연결된 조항이 없는 정책: {', '.join(none) if none else '없음'}", ""]
    return "\n".join(L_) + "\n"


def workbook(docs, reqs, basis):
    wb = Workbook()
    F = "맑은 고딕"
    head = PatternFill("solid", fgColor="1F3864")
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(ws, names, widths, rows):
        for i, n in enumerate(names, 1):
            c = ws.cell(1, i, n)
            c.font, c.fill, c.border, c.alignment = Font(name=F, bold=True, color="FFFFFF", size=10), head, box, Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[L(i)].width = widths[i - 1]
        for r, row in enumerate(rows, 2):
            for c, v in enumerate(row, 1):
                cell = ws.cell(r, c, v)
                cell.font, cell.alignment, cell.border = Font(name=F, size=9), wrap, box
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{L(len(names))}{len(rows) + 1}"

    ws = wb.active
    ws.title = "문서"
    sheet(ws, ["ID", "문서", "관할", "유형", "구속력", "상태", "주요 일자", "배포", "조항 수", "공식 페이지", "독립 검증"], [11, 50, 12, 16, 8, 12, 50, 18, 8, 50, 26],
          [[d["id"], d["title"], jur_label(d["jurisdiction"]), d["instrument_type"], "있음" if d["binding"] else "없음", STATUS_KO.get(d["status"], d["status"]), d["key_dates"], d["distribution"], d["provisions"], d["official_page_url"], review_label(d["review"])]
           for d in sorted(docs, key=lambda x: (x["jurisdiction_group"], x["id"]))])
    ws = wb.create_sheet("요구사항")
    sheet(ws, ["요구 ID", "문서", "관할", "조항", "주제", "강도", "구속력 문서", "대상", "기한", "요구 내용", "MINT 통제", "정책", "인용 문구"], [13, 10, 10, 18, 16, 8, 8, 10, 18, 60, 60, 20, 40],
          [[r["req_id"], r["doc_id"], jur_label(r["jurisdiction"]), r["ref"], THEME_KO.get(r["theme"], r["theme"]), STRENGTH_KO.get(r["strength"], r["strength"]), "있음" if r["binding_doc"] else "없음", r["actor"],
            r["timing_ko"], f"{r['topic_ko']} — {r['obligation_ko']}", r["gateway_control_ko"], ", ".join(r["pol_ids"]), r["anchor"]] for r in reqs])
    ws = wb.create_sheet("정책별 근거")
    sheet(ws, ["정책", "영역", "정책명", "연결 조항", "구속력 문서의 의무", "관할", "문서", "요구 ID"], [9, 14, 36, 8, 10, 28, 40, 60],
          [[b["policy_id"], b["area"], b["title"], b["provisions"], b["binding_mandatory"], ", ".join(jur_label(j) for j in b["jurisdictions"]), ", ".join(b["documents"]), ", ".join(b["requirements"])] for b in basis])
    juris = sorted({r["jurisdiction_group"] for r in reqs}, key=lambda j: (j != "KR", j))
    ws = wb.create_sheet("주제×관할")
    cnt = defaultdict(Counter)
    for r in reqs:
        cnt[r["theme"]][r["jurisdiction_group"]] += 1
    sheet(ws, ["주제"] + [jur_label(j) for j in juris] + ["합계"], [26] + [9] * (len(juris) + 1),
          [[THEME_KO[t]] + [cnt[t][j] or "" for j in juris] + [sum(cnt[t].values())] for t in THEME_KO if sum(cnt[t].values())])
    wb.save(OUT / "compliance-requirements.xlsx")


BEGIN, END = "<!-- compliance:begin -->", "<!-- compliance:end -->"


def papers_readme_block(docs):
    n_pub = sum(d["distribution"] == "public-under-license" for d in docs)
    lines = [BEGIN, "## 2026-10-17 해외·국내 AI 컴플라이언스 근거 문서 (L·M, %d건)" % len(docs), "",
             "AI 컴플라이언스 문서 모음(L, 12건)과 미국·EU·중국·캐나다·싱가포르·영국·일본·호주·인도·국제기구의 법·규제·정책·표준(M)을 공식 원문에서 확보해 조항 단위로 정리했다. "
             "조항별 요구사항·정책50 연결·관할×주제 집계·개발 계획은 [analysis/compliance](../analysis/compliance/README.md)에 있다. "
             "원문 재배포가 허용된 %d건만 `sources/compliance/`에 두고, 나머지 %d건은 git에서 제외한 `local-sources/compliance/`에만 보관하므로 다른 PC에서는 출처 URL과 해시로 다시 받아야 한다. "
             "이 자료는 법률 자문이 아니며 조항 정리는 읽은 범위 안에서만 유효하다." % (n_pub, len(docs) - n_pub), "",
             "| ID | 자료 | 관할 | 유형 | 상태 | 배포 | 조항 |", "| --- | --- | --- | --- | --- | --- | ---: |"]
    for d in sorted(docs, key=lambda x: (x["id"][:5], int(x["id"][5:]))):
        lines.append(f"| {d['id']} | [{esc(d['title'])}]({d['id']}.md) | {jur_label(d['jurisdiction'])} | {d['instrument_type']} | {STATUS_KO.get(d['status'], d['status'])} | "
                     f"{'공개' if d['distribution'] == 'public-under-license' else '로컬'} | {d['provisions']} |")
    lines += ["", END, ""]
    return "\n".join(lines)


def write_papers_readme(docs):
    path = WIKI / "papers/README.md"
    text = path.read_text(encoding="utf-8")
    block = papers_readme_block(docs)
    if BEGIN in text:
        head, rest = text.split(BEGIN, 1)
        tail = rest.split(END, 1)[1].lstrip("\n")
        text = head + block + "\n" + tail
    else:
        text = text.replace("## 업데이트", block + "\n## 업데이트", 1)
    path.write_text(text, encoding="utf-8", newline="\n")


def fill_block(text, name, body):
    begin, end = f"<!-- {name}:begin -->", f"<!-- {name}:end -->"
    head, rest = text.split(begin, 1)
    return head + begin + "\n" + body.strip("\n") + "\n" + end + rest.split(end, 1)[1]


def write_readme_blocks(docs, reqs, basis):
    path = OUT / "README.md"
    if not path.exists():
        return
    by = defaultdict(lambda: [0, 0, 0, 0])
    for d in docs:
        row = by[d["jurisdiction_group"]]
        row[0] += 1
        row[1] += d["distribution"] == "public-under-license"
        row[2] += d["provisions"]
    for r in reqs:
        if r["binding_mandatory"]:
            by[r["jurisdiction_group"]][3] += 1
    lines = [f"문서 {len(docs)}건(공개 원문 {sum(d['distribution'] == 'public-under-license' for d in docs)}건, 로컬 전용 {sum(d['distribution'] == 'local-only' for d in docs)}건), 조항 {len(reqs)}건, "
             f"구속력 있는 문서의 의무 조항 {sum(r['binding_mandatory'] for r in reqs)}건, 근거가 연결된 정책 {sum(1 for b in basis if b['provisions'])}/{len(basis)}건. "
             f"작성자와 다른 검증자가 전 조항을 원문과 대조한 문서는 {sum(1 for d in docs if d['review'])}/{len(docs)}건이다(문서별 결과는 근거 문서 목록의 '독립 검증' 열, 원자료는 `review-log.json`).", "",
             "| 관할 | 문서 | 공개 원문 | 조항 | 구속력 문서의 의무 조항 |", "| --- | ---: | ---: | ---: | ---: |"]
    for j in sorted(by, key=lambda j: (j != "KR", j)):
        a = by[j]
        lines.append(f"| {jur_label(j)} | {a[0]} | {a[1]} | {a[2]} | {a[3]} |")
    text = fill_block(path.read_text(encoding="utf-8"), "stats", "\n".join(lines))
    blocked = OUT / "acquisition-blocked.json"
    items = json.loads(blocked.read_text(encoding="utf-8"))["items"] if blocked.exists() else []
    if items:
        rows = [f"- **{esc(i['title'])}** ({jur_label(i['jurisdiction'])}, 예약 ID {i['reserved_id']}) — [공식 페이지]({i['official_page_url']}) · [PDF]({i['pdf_url']}). {esc(i['reason_ko'])} {esc(i['next_step_ko'])}" for i in items]
    else:
        rows = ["없음"]
    text = fill_block(text, "unavailable", "\n".join(rows))
    path.write_text(text, encoding="utf-8", newline="\n")


def write_license_record(notes):
    rows = [{"id": n["id"], "verified_on": TODAY, "evidence_url": n["official_page_url"], "license": n["license_notice"], "distribution": n["distribution"],
             "reason": n.get("distribution_reason", "")} for n in notes]
    write_json(WIKI / "sources/licenses/redistribution-2026-10-17.json", rows)


def main():
    notes, docs, reqs, pol = build()
    OUT.mkdir(parents=True, exist_ok=True)
    write_json(OUT / "requirements.json", {"generated_on": TODAY, "generator": "tools/build_compliance_matrix.py", "documents": docs, "requirements": reqs})
    write_text(OUT / "documents.md", md_documents(docs))
    write_text(OUT / "jurisdiction-theme.md", md_matrix(reqs))
    by_theme = defaultdict(list)
    for r in reqs:
        by_theme[r["theme"]].append(r)
    for th, rs in by_theme.items():
        write_text(OUT / "by-theme" / f"{th}.md", md_theme(th, rs))
    index = ["# 주제별 요구사항", "", "주제마다 관할별 조항을 모았다. 개수는 `전체(구속력 있는 문서의 의무 조항)`이다.", "", "| 주제 | 조항 | 의무(구속력 문서) |", "| --- | ---: | ---: |"]
    for th in THEME_KO:
        if th in by_theme:
            rs = by_theme[th]
            index.append(f"| [{THEME_KO[th]}]({th}.md) | {len(rs)} | {sum(r['binding_mandatory'] for r in rs)} |")
    write_text(OUT / "by-theme" / "README.md", "\n".join(index) + "\n")
    basis = pol_basis(reqs, pol)
    write_json(OUT / "pol-legal-basis.json", {"generated_on": TODAY, "policies": basis})
    write_text(OUT / "pol-legal-basis.md", md_pol(basis, {r["req_id"]: r for r in reqs}))
    workbook(docs, reqs, basis)
    write_papers_readme(docs)
    write_license_record(notes)
    write_readme_blocks(docs, reqs, basis)
    c = Counter(r["jurisdiction_group"] for r in reqs)
    print(f"documents {len(docs)}, requirements {len(reqs)}, binding mandatory {sum(r['binding_mandatory'] for r in reqs)}, policies with basis {sum(1 for b in basis if b['provisions'])}/{len(basis)}")
    print(dict(c))


if __name__ == "__main__":
    main()
